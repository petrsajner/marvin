"""Evaluation scripts: checkers, persistence, API endpoints and scoring hook."""
import copy
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from fastapi.testclient import TestClient

from harness import evals
from harness.application import ApplicationService
from harness.config import Config, load_config
from harness.session import Session


class EvalCheckerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        data = copy.deepcopy(load_config().data)
        data["agent"].update(workspace=None, autonomy="auto")
        self.cfg = Config(data, self.root)
        self.service = ApplicationService(self.cfg, llm_factory=lambda c: None,
                                          manage_model=False)

    def tearDown(self):
        self.service.close()
        self.service.models.wait(3)
        self.temp.cleanup()

    def _session(self, workspace: Path, script_id: str) -> Session:
        session = Session(self.cfg, system_prompt="SYS", workspace=str(workspace))
        session.meta["eval"] = script_id
        return session

    def test_code_add_checker_passes_on_correct_solution(self):
        ws = self.root / "ws1"
        ws.mkdir()
        (ws / "tests").mkdir()
        (ws / "textutils.py").write_text(
            "def word_frequency(text):\n    import string\n    out = {}\n"
            "    for raw in text.split():\n        word = raw.strip(string.punctuation).lower()\n"
            "        if word:\n            out[word] = out.get(word, 0) + 1\n    return out\n",
            encoding="utf-8")
        (ws / "tests" / "test_textutils.py").write_text(
            "import unittest\nfrom textutils import word_frequency\n"
            "class TestWordFrequency(unittest.TestCase):\n"
            "    def test_basic(self):\n"
            "        self.assertEqual(word_frequency('A a, b'), {'a': 2, 'b': 1})\n"
            "    def test_empty(self):\n"
            "        self.assertEqual(word_frequency(''), {})\n"
            "if __name__ == '__main__':\n    unittest.main()\n", encoding="utf-8")
        record = evals.score(self.cfg, self._session(ws, "code_add_function"), None)
        self.assertEqual(record["state"], "pass", record["detail"])

    def test_code_add_checker_fails_when_files_missing(self):
        ws = self.root / "ws2"
        ws.mkdir()
        record = evals.score(self.cfg, self._session(ws, "code_add_function"), None)
        self.assertEqual(record["state"], "fail")
        self.assertIn("missing files", record["detail"])

    def test_fix_bug_fixture_and_checker(self):
        ws = evals._make_workspace(self.cfg, "code_fix_bug")
        evals.EVALS["code_fix_bug"]["build_fixture"](ws)
        record = evals.score(self.cfg, self._session(ws, "code_fix_bug"), None)
        self.assertEqual(record["state"], "fail", "buggy fixture must fail")
        (ws / "calc.py").write_text(
            "def add_clamped(a, b, low=-100, high=100):\n"
            "    return max(low, min(high, a + b))\n", encoding="utf-8")
        record = evals.score(self.cfg, self._session(ws, "code_fix_bug"), None)
        self.assertEqual(record["state"], "pass", record["detail"])

    def test_document_checker(self):
        from docx import Document
        ws = self.root / "ws3"
        ws.mkdir()
        session = self._session(ws, "document_edit")
        record = evals.score(self.cfg, session, None)
        self.assertEqual(record["state"], "fail")
        doc = Document()
        for heading in evals._REQUIRED_HEADINGS:
            doc.add_heading(heading, level=1)
            doc.add_paragraph("Content paragraph with enough text to pass the length check. " * 2)
        doc.save(str(ws / "report.docx"))
        record = evals.score(self.cfg, session, None)
        self.assertEqual(record["state"], "pass", record["detail"])

    def test_memory_checker(self):
        ws = self.root / "ws4"
        ws.mkdir()
        session = self._session(ws, "chat_memory")
        record = evals.score(self.cfg, session, None)
        self.assertEqual(record["state"], "fail")
        (ws / "QWEN_MEMORY.md").write_text(
            "# Project memory\n- All measurements use metric units.\n", encoding="utf-8")
        record = evals.score(self.cfg, session, None)
        self.assertEqual(record["state"], "pass", record["detail"])

    def test_research_checker_reads_ledger(self):
        ws = self.root / "ws5"
        ws.mkdir()
        session = self._session(ws, "research_coverage")
        from harness.research import ResearchLedger
        ledger = ResearchLedger(session)
        ledger.begin("test question")
        ledger.record_source("https://one", "One", "x" * 300)
        ledger.record_source("https://two", "Two", "y" * 300)
        record = evals.score(self.cfg, session, None)
        self.assertEqual(record["state"], "fail")  # no synthesis yet
        run = ledger.current()
        run["synthesis"] = "[S1] says x. [S2] says y. " * 10
        run["citation_coverage"] = {"S1": True, "S2": True}
        run["status"] = "complete"
        ledger._save()
        record = evals.score(self.cfg, session, None)
        self.assertEqual(record["state"], "pass", record["detail"])

    def test_results_persist_across_instances(self):
        ws = self.root / "ws6"
        ws.mkdir()
        session = self._session(ws, "chat_memory")
        (ws / "QWEN_MEMORY.md").write_text("metric units fact\n", encoding="utf-8")
        evals.score(self.cfg, session, None)
        history = evals.read_history(self.cfg)
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0]["state"], "pass")
        # Scoring a non-eval session returns None and writes nothing.
        plain = Session(self.cfg, system_prompt="SYS", workspace=str(ws))
        self.assertIsNone(evals.score(self.cfg, plain, None))
        self.assertEqual(len(evals.read_history(self.cfg)), 1)


class EvalEndpointTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        data = copy.deepcopy(load_config().data)
        data["hardware"]["vram_gb"] = 32
        data["agent"].update(workspace=None, autonomy="auto")
        self.cfg = Config(data, self.root)
        self.service = ApplicationService(self.cfg, llm_factory=lambda c: None,
                                          manage_model=False)
        self.client = TestClient(create_app(self.cfg, service=self.app) if False else
                                 __import__("harness.web_api", fromlist=["create_app"])
                                 .create_app(self.cfg, service=self.service))

    def tearDown(self):
        self.service.close()
        self.service.models.wait(3)
        self.temp.cleanup()

    def test_catalog_lists_scripts_with_empty_history(self):
        catalog = self.client.get("/api/evals").json()
        self.assertEqual(len(catalog["scripts"]), 5)
        self.assertEqual(catalog["history"], [])
        self.assertTrue(all(s["last"] is None for s in catalog["scripts"]))

    def test_run_creates_session_with_meta_and_queues_job(self):
        with unittest.mock.patch.object(ApplicationService, "submit") as submit:
            response = self.client.post("/api/evals/code_fix_bug/run")
            response.raise_for_status()
            session_id = response.json()["session_id"]
            submit.assert_called_once()
            args, kwargs = submit.call_args
            self.assertEqual(args[0], session_id)
            self.assertIn("calc.py", args[1])
        session = self.service.session(session_id)
        self.assertEqual(session.meta.get("eval"), "code_fix_bug")
        self.assertEqual(session.meta.get("work_mode"), "development")
        self.assertEqual(
            self.client.post("/api/evals/nope/run").status_code, 400)


import unittest.mock  # noqa: E402  (used by EvalEndpointTests)

if __name__ == "__main__":
    unittest.main()
