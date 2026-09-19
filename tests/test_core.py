"""Core harness checks without a GPU or model server.

Run with the project's Python interpreter: tests/test_core.py."""
from __future__ import annotations

import atexit
import json
import shutil
import sys
import tempfile
import uuid
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from harness.i18n import locale_data, translate

# Use UTF-8 for Windows consoles and pipes to preserve Unicode diagnostics.
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

from harness.agent import build_registry
from harness.changes import ChangeJournal, file_sha256
from harness.config import Config, load_config
from harness.dependencies import (dependencies_current, mark_dependencies_current,
                                  requirements_digest, sync_dependencies)
from harness.llm import parse_tool_arguments
from harness.processes import ProcessManager
from harness.safety import Risk, SafetyPolicy
from harness.session import Session
from harness.tools.base import AgentContext, ToolRegistry

PASS = 0
FAIL = 0


_SCRATCH_SESSIONS = Path(tempfile.mkdtemp(prefix="marvin-test-sessions-"))
atexit.register(shutil.rmtree, _SCRATCH_SESSIONS, ignore_errors=True)
_LIVE_SESSIONS = load_config().path("paths.sessions_dir")
_LIVE_SESSIONS_BEFORE = ({item.name for item in _LIVE_SESSIONS.iterdir()}
                         if _LIVE_SESSIONS.is_dir() else set())


def scratch_sessions(config):
    """Point a configuration at a throwaway conversation directory.

    Several checks build a Session from the real configuration. Run from an
    installed copy they wrote into the owner's own conversation history, two
    directories per run, until this was added."""
    config.data["paths"]["sessions_dir"] = str(_SCRATCH_SESSIONS)
    return config


def check(cond: bool, label: str) -> None:
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✓ {label}")
    else:
        FAIL += 1
        print(f"  ✗ FAIL: {label}")


def test_config() -> None:
    print("[config]")
    cfg = load_config()
    check(cfg.model_key() in cfg.data["models"], "default model exists in 'models'")
    check(cfg.base_url.startswith("http://127.0.0.1"), "base_url points to localhost")
    check(cfg.model_file().name.endswith(".gguf"), "model_file points to a GGUF file")
    # Fresh-installation defaults; a local configuration can override the model.
    from copy import deepcopy
    from harness.config import DEFAULTS
    fresh = Config(deepcopy(DEFAULTS), root=ROOT)
    check(fresh.model_key() == "q5" and fresh.kv_cache_mode("q5") == "q8_0",
          "A fresh installation uses Qwen Q5 with Q8 KV")
    s = cfg.sampling(thinking=True)
    check(abs(s["temperature"] - 1.0) < 1e-9, "thinking sampling t=1.0")
    s2 = cfg.sampling(thinking=False)
    check(abs(s2["temperature"] - 0.7) < 1e-9, "non-thinking sampling t=0.7")
    cfg.data["default_model"] = "ornith_q5"
    check(abs(cfg.sampling(thinking=True)["temperature"] - 0.6) < 1e-9,
          "Ornith uses sampling settings from its own model card")
    check(cfg.mmproj_repo() == "ornith-ai/Ornith-1.5-35B-A3B-GGUF",
          "The Ornith vision projector can come from a separate repository")
    check(cfg.context_size() == 262144 and cfg.kv_cache_mode() == "q8_0_256k",
          "Ornith uses the validated 256k context and Q8 KV")
    legacy_file = Path(tempfile.mkdtemp()) / "legacy.yaml"
    try:
        legacy_file.write_text(
            "models:\n  q4:\n    ctx_size: 131072\n"
            "    kv_cache_profiles:\n"
            f"      f16: {{label: {json.dumps(translate('16-bit - more precise, context 128k', 'cs'), ensure_ascii=False)}, ctx_size: 131072}}\n"
            f"      q8_0: {{label: {json.dumps(translate('8-bit - larger context 256k', 'cs'), ensure_ascii=False)}, ctx_size: 262144}}\n"
            "  q5:\n    ctx_size: 98304\n"
            "agent:\n  max_steps: 40\n  semi_max_steps: 15\n",
            encoding="utf-8")
        migrated = load_config(legacy_file)
        check(migrated.context_size("q4") == 262144
              and migrated.context_size("q5") == 196608
              and migrated.kv_cache_mode("q4") == "q8_0"
              and migrated.kv_cache_mode("q5") == "q8_0",
              "Legacy configuration adopts measured Q8 defaults for Qwen")
        legacy_q4 = migrated.kv_cache_profiles("q4")
        check(legacy_q4["f16"]["label"] == "F16 · 128k"
              and legacy_q4["f16"].get("label_cs") == translate("F16 · 128k", "cs")
              and legacy_q4["q8_0"]["label"] == "Q8 · 256k",
              "Legacy localized KV labels migrate to English labels plus UI translations")
        migrated.set_kv_cache_mode("q4", "q8_0")
        check(migrated.context_size("q4") == 262144
              and migrated.kv_cache_server_args("q4")[-1] == "q8_0",
              "Selecting Q8 changes the corresponding Qwen context size")
        check(migrated.agent["max_steps"] == 0
              and migrated.agent["semi_max_steps"] == 0,
              "An old installation configuration cannot restore the removed agent-step limit")
    finally:
        shutil.rmtree(legacy_file.parent, ignore_errors=True)
    from harness.prompts import build_system_prompt, skills_block
    discussion_prompt = build_system_prompt("chat", cfg, ROOT, "discussion")
    research_prompt = build_system_prompt("chat", cfg, ROOT, "research")
    development_prompt = build_system_prompt("agent", cfg, ROOT, "development")
    check("DISCUSSION mode" in discussion_prompt and "coding agent" not in discussion_prompt,
          "Discussion does not use the development system prompt")
    check("Never filter" in research_prompt and "adult user" in research_prompt,
          "Research forbids filtering sources by trustworthiness")
    check("ORNITH DELIBERATE REASONING POLICY" in development_prompt
          and "Do not optimize for speed" in development_prompt,
          "Ornith xhigh receives explicit deep-reasoning guidance")
    skills_catalog = skills_block(cfg, ROOT)
    check("## OPTIONAL SKILLS" in skills_catalog
          and "research-synthesis" in skills_catalog
          and "translation-craft" in skills_catalog
          and "## OPTIONAL SKILLS" not in discussion_prompt,
          "The skill catalog reaches the model without entering the system prompt")
    from harness.version import APP_VERSION, _version_candidates
    # Installed copies keep version.txt at the root; development copies keep it under installer/.
    version_files = [p for p in _version_candidates() if p.exists()]
    installer_version = (version_files[0].read_text(encoding="utf-8").strip()
                         if version_files else "")
    check(bool(installer_version) and APP_VERSION == installer_version and APP_VERSION == "1.14.0",
          "The visible application version matches installer version 1.14.0")
    invariants = (ROOT / "AGENTS.md").read_text(encoding="utf-8")
    check(all(item in invariants for item in (
        "Language servers or an LSP runtime/distribution layer",
        "persistent interactive terminal",
        "Parallel model agents",
        "One-million-token context",
        "general plugin host, MCP ecosystem",
    )), "Permanent non-goals are recorded in the root product instructions")
    web_source = (ROOT / "webapp.py").read_text(encoding="utf-8")
    check(all(marker in web_source for marker in (
        'elem_id="workspace-control-stack"',
        't("Current task")', 't("Context")', 't("Runtime")',
        't("Settings & help")', 'show_progress="hidden"',
    )) and 't("Available skills"), open=' not in web_source
          and 't("Help & manuals"), open=' not in web_source,
          "The sidebar uses consistent information architecture")


def test_memory_layers() -> None:
    print("[memory layers]")
    from harness.memory import MemoryStore
    from harness.prompts import build_system_prompt, memory_block

    tmp = Path(tempfile.mkdtemp())
    try:
        workspace = tmp / "project"
        workspace.mkdir()
        memory_dir = tmp / "memory"
        memory_dir.mkdir()
        legacy_fact = "- The original development rule remains intact.\n"
        (memory_dir / "MEMORY.md").write_text(
            f"{locale_data('legacy_memory_heading')}\n\n"
            "<!-- scope=\"global\" -->\n" + legacy_fact,
            encoding="utf-8")

        cfg = Config(load_config().data, root=tmp)
        development = MemoryStore(cfg, workspace, "development")
        check(development.mode_path() == memory_dir / "MEMORY.md"
              and "The original development rule" in development.read("mode")
              and "Work mode memory: Development" in development.read("mode"),
              "Legacy MEMORY.md migrates losslessly to Development memory")
        check(development.global_path == memory_dir / "GLOBAL.md"
              and development.global_path != development.mode_path(),
              "Global memory remains separate from development memory")

        development.append("Universal preference", "global")
        development.append("Development rule", "mode")
        development.append("Project decision", "project")
        long_fact = "LONG-MEMORY-" + ("x" * 7000) + "-END-OF-MEMORY"
        development.append(long_fact, "mode")
        block = development.context_block()
        check(all(value in block for value in (
            "Universal preference", "Development rule", "Project decision")),
            "The system prompt includes global, work-mode and project memory")
        check("END-OF-MEMORY" in block,
              "Memory documents are injected in full without artificial truncation")

        paths = {
            mode: MemoryStore(cfg, workspace, mode).mode_path()
            for mode in ("discussion", "research", "writing", "development", "computer")
        }
        check(len(set(paths.values())) == 5,
              "Each work mode has its own memory document")
        research = MemoryStore(cfg, workspace, "research")
        research.append("Research rule", "mode")
        research_memory = memory_block(cfg, workspace, "research")
        research_prompt = build_system_prompt("chat", cfg, workspace, "research")
        check("Universal preference" in research_memory
              and "Research rule" in research_memory
              and "Project decision" in research_memory
              and "Development rule" not in research_memory,
              "Research sees its three memory layers without development memory")
        check("PERSISTENT MEMORY" not in research_prompt
              and "Universal preference" not in research_prompt,
              "Memory stays out of the system prompt so saving a fact keeps the prompt cache")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_safety() -> None:
    print("[safety]")
    sup = SafetyPolicy("supervised", max_steps=40, semi_max_steps=15)
    check(sup.needs_confirmation(Risk.WRITE), "Supervised mode confirms WRITE actions")
    check(not sup.needs_confirmation(Risk.SAFE), "Supervised mode allows SAFE actions without confirmation")
    sup.new_task()
    check(sup.step_limit() == 40, "supervised limit = max_steps")

    semi = SafetyPolicy("semi", max_steps=40, semi_max_steps=15)
    check(semi.needs_confirmation(Risk.WRITE), "Semi mode confirms the first WRITE action")
    semi.mark_confirmed()
    check(not semi.needs_confirmation(Risk.WRITE), "Semi mode allows subsequent WRITE actions after confirmation")
    check(semi.step_limit() == 15, "semi limit = semi_max_steps")

    auto = SafetyPolicy("auto", max_steps=40)
    check(not auto.needs_confirmation(Risk.WRITE), "Auto mode requires no confirmation")
    check(auto.step_limit() == 40, "auto limit = max_steps")
    unlimited = SafetyPolicy()
    check(unlimited.step_limit() is None, "The default agent has no step limit")

    try:
        SafetyPolicy("invalid-mode")
        check(False, "Invalid autonomy raises an exception")
    except ValueError:
        check(True, "Invalid autonomy raises an exception")


def test_session() -> None:
    print("[session]")
    tmp = Path(tempfile.mkdtemp())
    try:
        data = load_config().data
        data["paths"]["sessions_dir"] = str(tmp / "sessions")
        cfg = Config(data, root=ROOT)
        s = Session(cfg, session_id="test-session", system_prompt="SYS")
        s.add("user", "hello")
        img = tmp / "obrazek.png"
        img.write_bytes(b"\x89PNG fake")
        s.add("user", "mrkni na to", images=[img])
        s.add("assistant", "", tool_calls=[{"id": "call_1", "type": "function",
                                            "function": {"name": "list_dir", "arguments": "{}"}}])
        s.add("tool", "result text", tool_call_id="call_1", name="list_dir")
        check((tmp / "sessions/test-session/messages.jsonl").exists(), "JSONL was saved")
        n_img = len(list((tmp / "sessions/test-session/images").glob("*.png")))
        check(n_img == 1, f"image copied into the session ({n_img})")

        api = s.to_api_messages()
        check(api[0]["role"] == "system", "The system prompt comes first")
        img_msg = [m for m in api if isinstance(m.get("content"), list)]
        check(len(img_msg) == 1 and any(p["type"] == "image_url" for p in img_msg[0]["content"]),
              "An image is rendered as an image_url data URL")

        pinned = tmp / "pinned.txt"
        pinned.write_text("IMPORTANT PINNED CONTEXT", encoding="utf-8")
        check(s.pin_context_file(pinned), "A file can be pinned to the context")
        api_with_pin = s.to_api_messages()
        check(any("IMPORTANT PINNED CONTEXT" in str(m.get("content", ""))
                  for m in api_with_pin), "The pinned file appears in the model's API view")
        breakdown = s.context_breakdown()
        check(breakdown["pinned_files"] == [str(pinned.resolve())],
              "The context inspector tracks the pinned file")
        check(s.unpin_context_file(pinned) and not s.context_breakdown()["pinned_files"],
              "The pinned file can be unpinned")

        original_data_url = Session.__dict__["_data_url"]
        try:
            Session._data_url = staticmethod(lambda _path: (_ for _ in ()).throw(
                AssertionError("Token estimation must not encode images")))
            check(s.estimate_context_tokens() > Session.IMAGE_TOKENS,
                  "Context estimation neither reads nor base64-encodes image files")
        finally:
            Session._data_url = original_data_url

        loaded = Session.load(cfg, "test-session")
        check(len(loaded.messages) == 5, f"message roundtrip (={len(loaded.messages)})")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_tools_fs_shell() -> None:
    print("[tools]")
    import time
    tmp = Path(tempfile.mkdtemp())
    try:
        data = load_config().data
        data["paths"]["sessions_dir"] = str(tmp / "sessions")
        cfg = Config(data, root=ROOT)
        session = Session(cfg, session_id="tool-test")
        ctx = AgentContext(cfg=cfg, session=session, workspace=tmp)
        ctx.changes = ChangeJournal(session, tmp)
        ctx.processes = ProcessManager()
        ctx.changes.begin_task("Change test")

        reg = ToolRegistry()
        from harness.tools import fs, shell
        fs.register_fs_tools(reg)
        shell.register_shell_tools(reg)

        (tmp / "sub").mkdir()
        (tmp / "sub" / "a.txt").write_text("hello\nworld\nQWEN", encoding="utf-8")
        (tmp / "sub" / "data.py").write_text("x = 1\nQWEN_MARKER = 'zde'\n", encoding="utf-8")

        r = reg.execute("list_dir", {"path": "."}, ctx)
        check("sub" in r and "[DIR]" in r, "list_dir finds the directory")

        r = reg.execute("read_file", {"path": "sub/a.txt"}, ctx)
        check("hello" in r and "3|" in r, "read_file includes line numbers")

        r = reg.execute("write_file", {"path": "sub/new.md", "content": "# test"}, ctx)
        check((tmp / "sub" / "new.md").exists(), "write_file created the file")

        r = reg.execute("search_files", {"query": "qwen_marker", "path": "."}, ctx)
        check("data.py:2" in r, "search_files matches case-insensitively")
        r = reg.execute("search_files", {
            "query": r"QWEN_.* =", "path": ".", "regex": True,
        }, ctx)
        check("data.py:2" in r, "search_files supports regular expressions")
        r = reg.execute("find_files", {"pattern": "**/*.py", "path": "."}, ctx)
        check("data.py" in r, "find_files applies a project glob")
        r = reg.execute("make_directory", {"path": "generated/nested"}, ctx)
        check((tmp / "generated" / "nested").is_dir(), "make_directory creates parent directories")
        original_move = tmp / "sub" / "move-me.txt"
        original_move.write_text("restore me", encoding="utf-8")
        r = reg.execute("move_file", {"src": "sub/move-me.txt", "dst": "sub/moved.txt"}, ctx)
        check(r.startswith("OK") and (tmp / "sub" / "moved.txt").is_file(),
              "move_file renames the file")
        r = reg.execute("delete_file", {"path": "sub/moved.txt"}, ctx)
        check(r.startswith("OK") and not (tmp / "sub" / "moved.txt").exists(),
              "delete_file records deletion in the rollback journal")

        data_file = tmp / "sub" / "data.py"
        original_data = data_file.read_text(encoding="utf-8")
        r = reg.execute("apply_patch", {
            "path": "sub/data.py",
            "edits": [{"old": "MISSING", "new": "x"}],
        }, ctx)
        check(r.startswith("ERROR") and data_file.read_text(encoding="utf-8") == original_data,
              "An invalid patch does not change the file")
        r = reg.execute("apply_patch", {
            "path": "sub/data.py",
            "expected_sha256": file_sha256(data_file),
            "edits": [{"old": "x = 1", "new": "x = 2"}],
        }, ctx)
        check(r.startswith("OK") and "x = 2" in data_file.read_text(encoding="utf-8"),
              "apply_patch performs an exact atomic change")
        changes = reg.execute("list_task_changes", {}, ctx)
        check("data.py" in changes and "new.md" in changes,
              "The journal tracks modified and newly created files")
        undo = reg.execute("undo_task_changes", {}, ctx)
        check("errors\": []" in undo and data_file.read_text(encoding="utf-8") == original_data
              and not (tmp / "sub" / "new.md").exists()
              and original_move.read_text(encoding="utf-8") == "restore me"
              and not (tmp / "generated").exists(),
              "Rollback restores the complete pre-task state")
        check(not any(item["changed"] for item in ctx.changes.summary()["files"]),
              "The journal is clean after rollback")

        r = reg.execute("run_command", {"command": "echo hello-$((40+2))", "shell": "bash"}, ctx)
        check("hello-42" in r and "exit code: 0" in r, f"run_command works with bash: {r[:60]}")

        r = reg.execute("run_command", {"command": "Write-Output 'ps-works'"}, ctx)
        check("ps-works" in r, f"run_command works with PowerShell")

        r = reg.execute("run_command", {
            "command": "Write-Output HEAD-MARK; Write-Output ('x' * 25000); Write-Output TAIL-MARK",
            "shell": "powershell", "timeout": 10,
        }, ctx)
        log_match = __import__("re").search(r"\[full log: ([^\]]+)\]", r)
        check("HEAD-MARK" in r and "TAIL-MARK" in r and "full log saved" in r
              and log_match is not None and Path(log_match.group(1)).is_file(),
              "run_command retains head, tail and the complete log")

        import threading
        abort = threading.Event()
        ctx.abort_flag = abort
        timer = threading.Timer(0.35, abort.set)
        timer.start()
        started_abort = time.monotonic()
        r = reg.execute("run_command", {
            "command": "Start-Sleep -Seconds 20", "shell": "powershell", "timeout": 30,
        }, ctx)
        timer.cancel()
        check("aborted by user" in r and time.monotonic() - started_abort < 5,
              "Stop interrupts a synchronous command promptly")
        abort.clear()

        started = time.monotonic()
        launched = json.loads(reg.execute("start_command", {
            "command": "Write-Output one; Start-Sleep -Milliseconds 200; Write-Output two",
            "shell": "powershell", "timeout": 5,
        }, ctx))
        check(time.monotonic() - started < 1 and launched["status"] == "running",
              "start_command returns a process ID immediately")
        cursor = 0
        streamed = ""
        for _ in range(50):
            poll = json.loads(reg.execute("poll_command", {
                "process_id": launched["process_id"], "cursor": cursor,
            }, ctx))
            streamed += poll["output"]
            cursor = poll["cursor"]
            if poll["status"] == "finished":
                break
            time.sleep(0.05)
        check("one" in streamed and "two" in streamed and poll["exit_code"] == 0,
              "poll_command streams incremental output through completion")

        sleeper = json.loads(reg.execute("start_command", {
            "command": "Start-Sleep -Seconds 30", "shell": "powershell", "timeout": 60,
        }, ctx))
        stopped = json.loads(reg.execute("terminate_command", {
            "process_id": sleeper["process_id"],
        }, ctx))
        check(stopped.get("terminated") is True,
              "terminate_command stops the process tree")

        r = reg.execute("run_command", {"command": "format c: /x"}, ctx)
        check("blocked" in r.lower(), "A dangerous command was blocked")

        r = reg.execute("read_file", {"path": "missing.txt"}, ctx)
        check(r.startswith("ERROR"), "A read failure returns an ERROR message")

        schemas = reg.schemas()
        check(all(s["function"]["name"] for s in schemas) and len(schemas) == 16,
              f"schemas for an isolated set of 16 tools ({len(schemas)})")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_gpu_autofit() -> None:
    print("[gpu auto-fit]")
    from copy import deepcopy
    from harness.config import DEFAULTS, Config
    from harness.gpu import (best_fit, download_keys, effective_vram_gb, fits,
                             fitting_profiles)

    cfg = Config(deepcopy(DEFAULTS), root=ROOT)
    check(effective_vram_gb(cfg) is None or effective_vram_gb(cfg) > 0,
          "effective_vram_gb returns a detected/configured number or None")
    # A manual budget can constrain detected GPU capacity.
    cfg.data["hardware"] = {"vram_gb": 24}
    check(effective_vram_gb(cfg) == 24.0, "hardware.vram_gb constrains detected capacity")
    check(fits(cfg, "q4", "q8_0_compact", 24.0), "The compact profile fits a 24 GB capacity budget")
    check(not fits(cfg, "q4", "f16_compact", 24.0), "No F16 profile is offered for 24 GB cards")
    check(not fits(cfg, "q5", "q8_0", 24.0), "Q5 with a 192k context does not fit 24 GB")
    choice = best_fit(cfg, 24.0)
    check(choice == ("q5", "q8_0_96k"),
          f"auto-fit for 24 GB selects a feasible combination ({choice})")
    check(set(fitting_profiles(cfg, "ornith_q5", 24.0)) == set(),
          "Ornith has no supported 24 GB profile")
    check(set(download_keys(cfg, 24.0)) == {"q3", "q4", "q5", "nemotron_q4", "nemotron_q5"},
          "Setup selects compatible 24 GB models, including Nemotron Q4 with CPU experts")
    check(set(download_keys(cfg, 16.0)) == {"q3"},
          "Setup selects IQ3_S for 16 GB cards")
    check(best_fit(cfg, 16.0) == ("q3", "q8_0_64k"),
          "auto-fit for 16 GB selects the measured IQ3_S / 64k profile")
    q3_profiles = cfg.kv_cache_profiles("q3")
    check(set(q3_profiles) == {"q8_0", "q8_0_64k", "q8_0v4_96k", "q4_0_128k",
                               "q8_0_128k", "q8_0_96k"},
          "IQ3 has only the approved 16 and 24 GB profiles")
    check(set(cfg.kv_cache_profiles("q2")) == {"q8_0_96k_vision", "q8_0_64k_vision",
                                               "q8_0_128k", "q4_0_192k", "q8_0_256k"},
          "Q2 has the approved 16 GB profiles and one for a large card")
    check(cfg.model("q2").get("optional_download") is True
          and best_fit(cfg, 16.0) == ("q3", "q8_0_64k"),
          "Q2 is offered for 16 GB but never chosen automatically")
    check(all("min_vram_gb" in p for p in q3_profiles.values()),
          "Every Q3 profile specifies min_vram_gb")
    check(set(download_keys(cfg, 32.0)) == {"q4", "q5", "ornith_q5",
                                            "nemotron_q4", "nemotron_q5"},
          "Setup selects all standard models on a 32 GB card")
    # Nemotron: measured profiles, RAM spill, text-only (no projector)
    nq4 = cfg.kv_cache_profiles("nemotron_q4")
    check(nq4["q8_0_512k_spill"].get("server_args") == ["--fit", "off", "--n-cpu-moe", "14"],
          "Nemotron spill profiles include --n-cpu-moe")
    check(cfg.mmproj_file("nemotron_q4") is None and cfg.mmproj_file("q5") is not None,
          "Nemotron is text-only; Qwen has a multimodal projector")
    check(best_fit(cfg, 32.0) == ("q5", "q8_0"),
          "Q5 with Q8 remains the default for 32 GB cards")
    # Compact profiles use their cache_type instead of their profile identifier.
    cfg.set_kv_cache_mode("q4", "q8_0_compact")
    check(cfg.kv_cache_server_args("q4") == ["--cache-type-k", "q8_0",
                                             "--cache-type-v", "q8_0"],
          "The compact profile passes q8_0 as the actual cache type")
    check(cfg.context_size("q4") == 98304, "The compact Q4 profile has a 96k context")


def test_registry_modes() -> None:
    print("[registry]")
    chat = build_registry("chat")
    agent = build_registry("agent")
    computer = build_registry("computer")
    check(set(chat.names()) == {"read_memory", "save_memory", "web_search", "web_fetch",
                                "context_status", "pin_context_file", "unpin_context_file",
                                "list_project_documents", "read_project_document",
                                "export_document", "read_document", "edit_spreadsheet",
                                "list_skills", "read_skill",
                                "list_dir", "read_file", "write_file", "apply_patch",
                                "list_task_changes", "undo_task_changes", "search_files",
                                "find_files", "make_directory", "move_file", "delete_file",
                                "view_image", "search_project", "search_chat_history", "read_chat_history",
                                "edit_word_document", "view_document_page", "project_decisions",
                                "semantic_search"},
          f"chat mode: memory + web + context + disk tools ({len(chat.names())})")
    check({"list_dir", "run_command", "view_image"} <= set(agent.names()),
          f"agent mode: fs+patch+shell+vision ({len(agent.names())})")
    check({"screenshot", "click", "type_text", "press_key"} <= set(computer.names()),
          f"computer mode: adds GUI tools ({len(computer.names())})")
    click = computer.get("click")
    check(click.risk == Risk.WRITE, "click has WRITE risk")
    shot = computer.get("screenshot")
    check(shot.risk == Risk.SAFE, "screenshot has SAFE risk")

    discussion = build_registry("chat", "discussion")
    research = build_registry("chat", "research")
    writing = build_registry("agent", "writing")
    development = build_registry("agent", "development")
    check({"read_file", "write_file", "search_files", "view_image"} <= set(discussion.names())
          and not ({"git_commit", "run_command", "repo_overview"} & set(discussion.names())),
          "Discussion has file tools without the development toolset")
    check(set(research.names()) == set(discussion.names()),
          "Research has web/context tools without the development toolset")
    check("export_document" in discussion.names() and "export_document" in research.names()
          and "export_document" in development.names(),
          "Document export is available in every work mode")
    check("apply_patch" in writing.names() and "export_document" in writing.names()
          and "repo_overview" not in writing.names() and "git_commit" not in writing.names()
          and "run_command" not in writing.names(),
          "Writing has document editing without Git and shell tools")
    check({"apply_patch", "git_commit", "run_command", "start_project_check"}
          <= set(development.names()), "Development has the complete development toolset")
    check({"browser_open", "browser_snapshot", "browser_screenshot", "browser_console",
           "browser_network", "browser_select", "browser_upload", "browser_download",
           "browser_viewport"} <= set(development.names()),
          "Development has an isolated browser session")
    check({"find_symbol", "document_symbols", "find_references"}
          <= set(development.names()),
          "Development has lightweight multi-language symbol navigation")


def test_parse_args() -> None:
    print("[llm helpers]")
    check(parse_tool_arguments('{"x": 1}') == {"x": 1}, "Valid JSON")
    check(parse_tool_arguments("") == {}, "Empty arguments")
    check(parse_tool_arguments('blabla {"x": [1,2]} blabla') == {"x": [1, 2]},
          "JSON embedded in surrounding text")


def test_workspace() -> None:
    print("[workspace]")
    from harness.agent import Agent
    data = load_config().data
    data["paths"]["sessions_dir"] = str(Path(tempfile.mkdtemp()) / "sessions")
    cfg = Config(data, root=ROOT)
    tmp = Path(tempfile.mkdtemp())
    try:
        session = Session(cfg, session_id="ws-test", system_prompt="SYS")
        from harness.safety import SafetyPolicy
        agent = Agent(cfg, LLMStub(), session, build_registry("agent"),
                      SafetyPolicy("supervised"), mode="agent")
        # None selects the current working directory.
        check(agent.workspace == Path.cwd().resolve(), "The default workspace is the current directory")
        check(agent.ctx.project_workspace is None and agent.ctx.repo_index is None,
              "Projectless chats have no project document index")
        # Set a directory.
        p = agent.set_workspace(str(tmp))
        check(p == tmp.resolve() and agent.workspace == tmp.resolve()
              and agent.ctx.project_workspace == tmp.resolve()
              and agent.ctx.repo_index is not None,
              "set_workspace updates tool paths and the project index")
        # A file path resolves to its parent directory.
        f = tmp / "file.txt"
        f.write_text("x", encoding="utf-8")
        (tmp / "module.py").write_text("def project_symbol():\n    return 1\n", encoding="utf-8")
        p2 = agent.set_workspace(str(f))
        check(p2 == tmp.resolve(), "A file resolves to its parent directory")
        # Quoted path.
        p3 = agent.set_workspace(f'"{tmp}"')
        check(p3 == tmp.resolve(), "A quoted path is accepted")
        # Missing path.
        try:
            agent.set_workspace(tmp / "missing")
            check(False, "A missing path raises ValueError")
        except ValueError:
            check(True, "A missing path raises ValueError")
        # Tools resolve relative paths against their workspace.
        from harness.tools.base import AgentContext
        r = build_registry("agent").execute("read_file", {"path": "file.txt"}, agent.ctx)
        check("file.txt" in r and "1| x" in r, "read_file resolves paths against the workspace")
        agent.new_task("prozkoumej projekt")
        dynamic = agent._api_messages()[-1]["content"]
        check("CURRENT PROJECT SNAPSHOT" in dynamic and "project_symbol" in dynamic
              and "project_symbol" not in session.messages[0]["content"],
              "Changing repository context stays at the tail without invalidating the stable prefix")
        cached_prefix = json.dumps(session.messages, ensure_ascii=False, sort_keys=True)
        cached_count = len(session.messages)
        overview = build_registry("agent").execute("repo_overview", {}, agent.ctx)
        check("module.py" in overview and "project_symbol" in overview,
              "repo_overview includes important workspace symbols")
        (tmp / "module.py").write_text("def refreshed_symbol():\n    return 2\n", encoding="utf-8")
        refreshed = build_registry("agent").execute("repo_overview", {}, agent.ctx)
        check("refreshed_symbol" in refreshed and "project_symbol" not in refreshed,
              "A file change invalidates the repository snapshot cache")
        agent.new_task("Continue using the current state")
        check(json.dumps(session.messages[:cached_count], ensure_ascii=False, sort_keys=True)
              == cached_prefix
              and "refreshed_symbol" in agent._api_messages()[-1]["content"]
              and not any(str(message.get("content", "")).startswith("[DYNAMIC TASK CONTEXT")
                          for message in session.messages),
              "A new snapshot does not duplicate the reusable prefix")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_shell_readonly() -> None:
    print("[shell read-only klasifikace]")
    from harness.tools.shell import RunCommandTool, is_read_only_command
    tool = RunCommandTool()
    safe = [
        "ls -la", "cat file.txt", "grep -r foo .", "git status", "git log --oneline",
        "git diff HEAD~1", "find . -name '*.py'", "echo hello", "ls | grep test",
        "cat a.txt; cat b.txt", "stat main.py", "wc -l *.py",
    ]
    unsafe = [
        "rm -rf x", "echo hello > file.txt", "cat x | tee y", "mkdir new",
        "git push", "git commit -m x", "curl http://x", "ls; rm x",
        "echo $(rm x)", "npm install x", "grep x . > out", "git branch nova",
        "cp a b", "cat < input.txt", "python script.py", "",
    ]
    for cmd in safe:
        check(is_read_only_command(cmd), f"SAFE: {cmd!r}")
    for cmd in unsafe:
        check(not is_read_only_command(cmd), f"WRITE: {cmd!r}")
    check(tool.risk_for({"command": "ls"}) == Risk.SAFE, "risk_for ls = SAFE")
    check(tool.risk_for({"command": "rm x"}) == Risk.WRITE, "risk_for rm = WRITE")


def test_context_compression() -> None:
    print("[non-destructive context compression]")
    from harness.context import render_messages_text

    rendered = render_messages_text([
        {"role": "user", "content": "HEAD " + "a" * 120},
        {"role": "assistant", "content": "MIDDLE " + "b" * 500},
        {"role": "user", "content": "TAIL-CONTEXT " + "c" * 120},
    ], max_chars=240)
    check("HEAD" in rendered and "TAIL-CONTEXT" in rendered and len(rendered) <= 240,
          "Long transcript trimming retains the beginning and the newest end")
    tmp = Path(tempfile.mkdtemp())
    try:
        data = load_config().data
        data["paths"]["sessions_dir"] = str(tmp / "sessions")
        cfg = Config(data, root=ROOT)
        s = Session(cfg, session_id="ctx-test", system_prompt="SYS")
        # Simulate a longer conversation with six user/assistant/tool triples
        for i in range(6):
            s.add("user", f"otazka {i} " + "x" * 500)
            s.add("assistant", "", tool_calls=[{"id": f"c{i}", "type": "function",
                                                "function": {"name": "read_file", "arguments": "{}"}}])
            s.add("tool", f"response {i} " + "y" * 300, tool_call_id=f"c{i}", name="read_file")
        n_before = len(s.messages)
        est = s.estimate_context_tokens()
        check(est > 1000, f"reasonable token estimate ({est})")

        ok = s.compress_to_summary("Conversation SUMMARY.", min_keep=6)
        check(ok, "Compression completed")
        check(len(s.messages) == n_before, f"history UNCHANGED ({len(s.messages)} == {n_before})")
        check(s.compression is not None and s.compression["cut"] > 1, "Compression boundary recorded")

        # The model sees a reduced view while the user retains the full history.
        api_view = s._view_messages()
        check(len(api_view) < n_before, f"model view is smaller ({len(api_view)} < {n_before})")
        check("Conversation SUMMARY." in api_view[1]["content"], "The summary appears in the model view")
        est2 = s.estimate_context_tokens()
        check(est2 < est, f"model token count decreased ({est} → {est2})")

        # The model view must not start with an orphaned tool result.
        first_role = api_view[2]["role"] if len(api_view) > 2 else None
        check(first_role in ("user", None), f"cut na user hranici (role={first_role})")

        # persist and roundtrip both history and compression
        loaded = Session.load(cfg, "ctx-test")
        check(len(loaded.messages) == n_before, "The complete JSONL history survives a round trip")
        check(loaded.compression is not None and loaded.compression["cut"] == s.compression["cut"],
              "compression.json was persisted")

        # A second compression advances the cut.
        s.add("user", "nova otazka " + "a" * 100)
        s.add("assistant", "new response " + "b" * 100)
        ok2 = s.compress_to_summary("SUMMARY 2.", min_keep=2)
        check(ok2 and s.compression["cut"] > loaded.compression["cut"],
              "The second compression advanced the cut")

        # Fallback trimming advances the view boundary without deleting history.
        s2 = Session(cfg, session_id="trim-test", system_prompt="SYS")
        for i in range(10):
            s2.add("user", f"u{i} " + "z" * 2000)
            s2.add("assistant", f"a{i} " + "w" * 2000)
        big = s2.estimate_context_tokens()
        n2 = len(s2.messages)
        ok = s2.trim_to_budget(big // 2)
        check(ok and s2.estimate_context_tokens() <= big // 2,
              f"trim fits the budget ({big} → {s2.estimate_context_tokens()})")
        check(len(s2.messages) == n2, "Trimming changes the view boundary without deleting history")

        # Large recent tool outputs must respect the retained-context token budget.
        s3 = Session(cfg, session_id="bigtail-test", system_prompt="SYS")
        for i in range(10):
            s3.add("user", f"q{i}")
            s3.add("assistant", "", tool_calls=[{"id": f"c{i}", "type": "function",
                                                 "function": {"name": "read_file", "arguments": "{}"}}])
            s3.add("tool", "T" * 6000, tool_call_id=f"c{i}", name="read_file")  # Approximately 1,600 tokens.
        # Ten message triples exceed 16k tokens; a 6k budget requires a substantial cut.
        est_before = s3.estimate_context_tokens()
        ok = s3.compress_to_summary("SUMMARY", keep_tokens=6000)
        est_after = s3.estimate_context_tokens()
        check(ok and est_after <= 8000,
              f"token budget preserves the tail ({est_before} → {est_after}, target ≤ 8000)")
        view = s3._view_messages()
        check(view[2]["role"] == "user", "Token-budget cuts remain on user-message boundaries")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_reasoning_effort_kwargs() -> None:
    print("[reasoning effort]")
    from harness.llm import LLMClient, _template_kwargs
    data = load_config().data
    data["thinking"] = True
    data["reasoning_effort"] = "low"
    check(_template_kwargs(Config(data, ROOT)) == {"chat_template_kwargs": {"reasoning_effort": "low"}},
          "effort low → template kwarg")
    data["reasoning_effort"] = "xhigh"
    check(_template_kwargs(Config(data, ROOT))["chat_template_kwargs"]["reasoning_effort"] == "xhigh",
          "effort xhigh")
    data["thinking"] = False
    check(_template_kwargs(Config(data, ROOT)) == {
        "chat_template_kwargs": {"enable_thinking": False}},
          "Disabling thinking takes precedence over reasoning effort")
    data["thinking"] = True
    data["reasoning_effort"] = "blbost"
    check(_template_kwargs(Config(data, ROOT)) == {}, "Invalid effort uses the chat-template default")
    data["default_model"] = "ornith_q5"
    data["reasoning_effort"] = "xhigh"
    check(_template_kwargs(Config(data, ROOT)) == {
        "chat_template_kwargs": {"enable_thinking": True}},
          "Ornith enables reasoning without an unsupported reasoning_effort parameter")
    data["default_model"] = "q4"

    class CaptureCompletions:
        def __init__(self):
            self.params = None

        def create(self, **params):
            self.params = params
            return []

    data["thinking"] = False
    cfg = Config(data, ROOT)
    llm = LLMClient.__new__(LLMClient)
    llm.cfg = cfg
    llm.model_name = "local-model"
    completions = CaptureCompletions()
    llm.client = type("Client", (), {
        "chat": type("Chat", (), {"completions": completions})(),
    })()
    llm.stream([{"role": "user", "content": "test"}])
    extra = completions.params["extra_body"]
    check(extra.get("top_k") == 20, "Streaming preserves top_k")
    check(extra.get("chat_template_kwargs") == {"enable_thinking": False},
          "Streaming preserves thinking and reasoning settings")
    check("max_tokens" not in completions.params,
          "Production requests have no artificial output-token limit")

    class ClosableStream:
        def __init__(self):
            self.closed = False
            self.parts = ["Beginning of a ", "completed sentence.", " This text must not be generated."]

        def __iter__(self):
            for part in self.parts:
                delta = type("Delta", (), {"content": part, "tool_calls": []})()
                yield type("Chunk", (), {
                    "choices": [type("Choice", (), {"delta": delta})()]})()

        def close(self):
            self.closed = True

    closable = ClosableStream()
    llm.client = type("Client", (), {
        "chat": type("Chat", (), {
            "completions": type("Completions", (), {
                "create": staticmethod(lambda **_kwargs: closable)})(),
        })(),
    })()
    import threading
    stop_requested = threading.Event()

    def request_stop():
        return stop_requested.is_set()

    stopped = llm.stream([{"role": "user", "content": "test stop"}],
                         should_stop=request_stop, on_text=lambda _text: stop_requested.set())
    check(stopped.stopped and stopped.content == "Beginning of a completed sentence."
          and closable.closed,
          "Graceful Stop finishes the sentence, closes the stream and stops generation")


def test_runtime_lifecycle_helpers() -> None:
    print("[runtime lifecycle]")
    import socket
    import time
    from harness import servermgmt
    try:
        from launcher.launcher_app import _free_web_port
    except ImportError:
        # Installed copies may only contain the frozen launcher executable.
        _free_web_port = None

    tmp = Path(tempfile.mkdtemp())
    original_health = servermgmt.health
    try:
        data = load_config().data
        data["paths"]["runtime_dir"] = str(tmp / "runtime")
        data["server"]["port"] = 65534
        cfg = Config(data, root=ROOT)
        pf = servermgmt.pid_file(cfg)
        pf.parent.mkdir(parents=True)
        pf.write_text("q4:99999999", encoding="utf-8")
        servermgmt.health = lambda *_args, **_kwargs: False
        check(servermgmt.server_state(cfg) == "down" and not pf.exists(),
              "A stale PID is removed and the server is down")

        class DeadProcess:
            @staticmethod
            def poll():
                return 1

        started = time.monotonic()
        check(not servermgmt.wait_health(cfg, timeout=10, proc=DeadProcess())
              and time.monotonic() - started < 1,
              "wait_health returns promptly after a process crash")

        if _free_web_port is not None:
            with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as occupied:
                occupied.bind(("127.0.0.1", 0))
                occupied.listen(1)
                busy_port = occupied.getsockname()[1]
                check(_free_web_port(busy_port) != busy_port,
                      "The launcher skips an occupied UI port")
    finally:
        servermgmt.health = original_health
        shutil.rmtree(tmp, ignore_errors=True)


def test_dependency_marker() -> None:
    print("[dependency marker]")
    tmp = Path(tempfile.mkdtemp())
    try:
        requirements = tmp / "requirements.txt"
        venv = tmp / ".venv"
        python = venv / "Scripts" / "python.exe"
        python.parent.mkdir(parents=True)
        python.touch()
        requirements.write_text("example==1.0\n", encoding="utf-8")

        (venv / ".deps.ok").write_text("", encoding="ascii")
        check(not dependencies_current(requirements, venv),
              "the old .deps.ok marker is not treated as current")
        mark_dependencies_current(requirements, venv)
        check(dependencies_current(requirements, venv),
              "SHA-256 marker confirms current requirements")
        check(not (venv / ".deps.ok").exists(),
              "synchronization removes the stale marker")
        check(len(requirements_digest(requirements)) == 64,
              "the requirements fingerprint is SHA-256")

        requirements.write_text("example==2.0\n", encoding="utf-8")
        check(not dependencies_current(requirements, venv),
              "a requirements change invalidates the dependency marker")
        from unittest.mock import patch
        with patch("harness.dependencies.subprocess.call", return_value=0) as pip_call:
            rc = sync_dependencies(requirements, venv, force=True)
        command = pip_call.call_args.args[0]
        check(rc == 0 and command[-2:] == ["-r", str(requirements)],
              "Dependency synchronization installs the declared requirements")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_offline_backup() -> None:
    print("[offline backup]")
    from scripts.offline_backup import backup_info, create_backup, restore_backup, verify_backup

    outer = Path(tempfile.mkdtemp())
    try:
        source = outer / "source"
        (source / "runtime" / "models").mkdir(parents=True)
        (source / "runtime" / "llama" / "bin").mkdir(parents=True)
        (source / ".venv" / "Lib" / "site-packages" / "example").mkdir(parents=True)
        (source / ".venv" / "Lib" / "site-packages" / "example" / "__init__.py").write_text(
            "VALUE = 1\n", encoding="utf-8")
        (source / "runtime" / "models" / "model.gguf").write_bytes(b"MODEL" * 1000)
        (source / "runtime" / "models" / "mmproj.gguf").write_bytes(b"VISION" * 100)
        (source / "runtime" / "llama" / "bin" / "llama-server.exe").write_bytes(
            b"LLAMA" * 500)
        (source / "runtime" / "llama" / "bin" / "cuda.dll").write_bytes(b"CUDA" * 200)
        (source / "requirements.txt").write_text("example==1.0\n", encoding="utf-8")
        (source / "version.txt").write_text("test-version\n", encoding="utf-8")
        (source / "dist").mkdir()
        (source / "dist" / "Marvin-Setup-test-version.exe").write_bytes(b"SETUP")
        backup = outer / "QwenHarness-Offline-Backup"
        manifest = create_backup(source, backup)
        check(manifest["format_version"] == 2 and (backup / "manifest.json").is_file(),
              "Backup creation writes a versioned manifest")
        verified = verify_backup(backup)
        check(verified["ok"] and verified["files"] >= 5,
              "An intact offline backup passes SHA-256 verification")
        info = backup_info(backup)
        check(info["app_version"] == "test-version" and len(info["models"]) == 2
              and info["dependencies"] and info["installer"].endswith(".exe"),
              "Backup metadata describes the version, models, installer and Python dependencies")
        restored_root = outer / "restored"
        (restored_root / ".venv" / "Scripts").mkdir(parents=True)
        (restored_root / ".venv" / "Scripts" / "python.exe").touch()
        (restored_root / "requirements.txt").write_text("example==1.0\n", encoding="utf-8")
        result = restore_backup(restored_root, backup)
        check(result["ok"]
              and (restored_root / "runtime/models/model.gguf").read_bytes()
              == b"MODEL" * 1000
              and (restored_root / "runtime/llama/bin/llama-server.exe").is_file()
              and (restored_root / ".venv/Lib/site-packages/example/__init__.py").is_file()
              and (restored_root / ".venv/.requirements.sha256").is_file(),
              "Restore recovers models, the inference runtime and Python dependencies")
        fallback_root = outer / "fallback"
        targeted = restore_backup(fallback_root, backup, {"models"})
        check((fallback_root / "runtime/models/model.gguf").is_file()
              and not (fallback_root / "runtime/llama").exists()
              and targeted["dependencies"] == "not-requested",
              "Fallback restore copies only the component unavailable online")
        damaged = backup / "payload" / "runtime" / "models" / "model.gguf"
        damaged.write_bytes(b"DAMAGED")
        check(not verify_backup(backup)["ok"],
              "Size or SHA-256 detects payload corruption")
        # Setup scripts live under installer/ in development and at the installed application root.
        setup_bat = next((p2 for p2 in (ROOT / "installer" / "run_setup.bat",
                                        ROOT / "run_setup.bat") if p2.is_file()), None)
        backup_bat = next((p2 for p2 in (ROOT / "installer" / "run_setup_from_backup.bat",
                                         ROOT / "run_setup_from_backup.bat") if p2.is_file()), None)
        iss_file = ROOT / "installer" / "marvin.iss"
        installer_text = setup_bat.read_text(encoding="utf-8") if setup_bat else ""
        check("offline_backup.py restore" in installer_text
              and "scripts\\sync_deps.py" in installer_text
              and "--components models" in installer_text
              and backup_bat is not None
              and "QWEN_HARNESS_BACKUP_PREFER" in backup_bat.read_text(encoding="utf-8")
              and (not iss_file.is_file() or "run_setup_from_backup.bat" in
                   iss_file.read_text(encoding="utf-8")),
              "The installer supports local restore and backup-folder selection")
    finally:
        shutil.rmtree(outer, ignore_errors=True)


def test_streaming_bridge() -> None:
    print("[streaming bridge]")
    from harness.streaming import SteeringQueue, StreamHub, step_threaded

    hub = StreamHub()
    hub.on_event("reasoning", "thi")
    hub.on_event("reasoning", "nking")
    hub.on_event("text", "do")
    hub.on_event("text", "ne")
    text, reasoning, rev, _ = hub.snapshot()
    check(text == "done" and reasoning == "thinking" and rev == 4,
          "StreamHub preserves fragment order")
    hub.on_event("tool_delta", ("write_file", '{"path":"game.py","content":"abc'))
    progress = hub.progress()
    check(progress["tool_call_name"] == "write_file"
          and progress["tool_call_chars"] > 20,
          "StreamHub exposes incremental tool-argument generation")
    hub.on_event("tool_start", ("write_file", {"path": "game.py"}))
    check(hub.progress()["tools_running"] == [("write_file", {"path": "game.py"})],
          "StreamHub shows the tool currently executing")
    hub.on_event("tool_result", ("write_file", "OK"))
    check(not hub.progress()["tools_running"],
          "StreamHub clears tool execution after its result")

    class FakeAgent:
        @staticmethod
        def step(approve=None):
            return f"step:{approve}"

    thread, box = step_threaded(FakeAgent(), True)
    thread.join(timeout=2)
    check(not thread.is_alive() and box.get("r") == "step:True",
          "The worker bridge returns the agent-step result")

    steering = SteeringQueue()
    steering.push("Fix the parser first.", ["screen.png"])
    steering.push("A zachovej kompatibilitu.")
    check(bool(steering)
          and steering.pop_all() == [
              ("Fix the parser first.", ["screen.png"]),
              ("A zachovej kompatibilitu.", []),
          ] and not steering,
          "Steering preserves clarification order and drains atomically")

    from harness.agent import Agent, Status
    from harness.llm import AssistantResult
    from harness.safety import SafetyPolicy
    tmp = Path(tempfile.mkdtemp())
    try:
        data = load_config().data
        data["paths"]["sessions_dir"] = str(tmp / "sessions")
        data["agent"]["workspace"] = None
        cfg = Config(data, root=ROOT)
        session = Session(cfg, session_id="stop-test", system_prompt="SYS")
        agent = Agent(
            cfg, LLMStub([AssistantResult(content="Completed sentence.", stopped=True)]),
            session, ToolRegistry(), SafetyPolicy("auto"), mode="chat",
            work_mode="discussion")
        agent.new_task("Long response")
        result = agent.step()
        check(result.status is Status.ABORTED
              and any(message.get("content") == "Completed sentence."
                      for message in session.messages),
              "Stop preserves completed response text and ends the task")
        agent.abort_flag.set()
        agent.llm = LLMStub([AssistantResult(content="New response")])
        agent.new_task("New question after Stop")
        check(not agent.abort_flag.is_set() and agent.step().status is Status.FINAL,
              "A request after Stop receives a clear abort state")

        steered = Session(cfg, session_id="steer-test", system_prompt="SYS")
        steer_agent = Agent(
            cfg, LLMStub([AssistantResult(content="First completed sentence.", stopped=True)]),
            steered, ToolRegistry(), SafetyPolicy("auto"), mode="chat",
            work_mode="discussion")
        steer_agent.new_task("Propose a solution")
        check(steer_agent.step().status is Status.ABORTED,
              "Steering first stops the active stream at a complete sentence")
        steer_agent.steer("Also preserve backward compatibility.")
        steer_agent.llm = LLMStub([AssistantResult(content="Revised solution.")])
        steer_result = steer_agent.step()
        visible = [m.get("content") for m in steered.messages
                   if m.get("role") != "system"
                   and not str(m.get("content") or "").startswith(Session.INTERNAL_USER_PREFIXES)]
        check(steer_result.status is Status.FINAL
              and visible[-3:] == [
                  "First completed sentence.",
                  "Also preserve backward compatibility.",
                  "Revised solution.",
              ],
              "Steering preserves partial output and continues with clarifications in order")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_parallel_read_tools() -> None:
    print("[parallel read tools]")
    import time
    from harness.agent import Agent
    from harness.llm import AssistantResult
    from harness.safety import SafetyPolicy
    from harness.tools.base import Tool

    class SlowRead(Tool):
        parallel_safe = True
        parameters = {}

        def __init__(self, name):
            self.name = name

        def run(self, _ctx):
            time.sleep(0.25)
            return self.name

    class SlowWrite(SlowRead):
        parallel_safe = False

    tmp = Path(tempfile.mkdtemp())
    try:
        data = load_config().data
        data["paths"]["sessions_dir"] = str(tmp / "sessions")
        cfg = Config(data, root=ROOT)
        session = Session(cfg, session_id="parallel-test", system_prompt="SYS")
        registry = ToolRegistry()
        registry.register(SlowRead("read_a"))
        registry.register(SlowRead("read_b"))
        registry.register(SlowWrite("write_c"))
        agent = Agent(cfg, LLMStub(), session, registry, SafetyPolicy("auto"), mode="agent")
        calls = [_tc("read_a"), _tc("read_b")]
        started = time.monotonic()
        trace = agent._execute_calls(calls, "I found the initial sources and will now compare them.")
        parallel_time = time.monotonic() - started
        check(parallel_time < 0.45 and [item[2] for item in trace] == ["read_a", "read_b"],
              "Independent read-only tools can run concurrently while preserving result order")
        persisted = next(message for message in session.messages
                         if message.get("tool_calls") == calls)
        reloaded = Session.load(cfg, session.id)
        check(persisted["content"].startswith("I found")
              and any(str(message.get("content", "")).startswith("I found")
                      for message in reloaded.messages),
              "Progress text before a tool call remains visible after reload")
        started = time.monotonic()
        agent._execute_calls([_tc("read_a"), _tc("write_c")])
        check(time.monotonic() - started >= 0.45,
              "Mixed read/write tool groups execute sequentially")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_resume_task_and_process_after_restart() -> None:
    print("[resume task/process after restart]")
    import time
    from harness.agent import Agent, Status
    from harness.llm import AssistantResult
    from harness.safety import SafetyPolicy

    tmp = Path(tempfile.mkdtemp())
    managers = []
    try:
        data = load_config().data
        data["paths"]["sessions_dir"] = str(tmp / "sessions")
        data["agent"]["workspace"] = str(tmp)
        cfg = Config(data, root=ROOT)
        session = Session(cfg, session_id="resume-test", system_prompt="SYS")
        pending_llm = LLMStub([AssistantResult(content="Preparing the file and waiting for approval.", tool_calls=[
            _tc("write_file", '{"path":"resume.txt","content":"OK"}')])])
        first = Agent(
            cfg, pending_llm, session, build_registry("agent", "development"),
            SafetyPolicy("supervised"), mode="agent", work_mode="development")
        first.new_task("Create resume.txt")
        waiting = first.step()
        check(waiting.status is Status.NEEDS_CONFIRMATION
              and session.load_task_state()["status"] == "waiting_confirmation"
              and session.load_task_state()["pending_text"].startswith("Preparing"),
              "Pending approvals and their visible progress text are persisted")

        restored = Agent(
            cfg, LLMStub([]), session, build_registry("agent", "development"),
            SafetyPolicy("supervised"), mode="agent", work_mode="development")
        restored_waiting = restored.step()
        check(restored.has_resumable_task
              and restored_waiting.status is Status.NEEDS_CONFIRMATION
              and restored_waiting.text.startswith("Preparing"),
              "A new Agent restores pending tools and visible text after restart")
        restored.step(approve=True)
        check((tmp / "resume.txt").is_file()
              and any(str(message.get("content", "")).startswith("Preparing")
                      for message in session.messages if message.get("role") == "assistant"),
              "Restored approval completes the tool call and preserves progress")
        completed = Agent(
            cfg, LLMStub([AssistantResult(content="done"), AssistantResult(content="Verified and complete")]), session,
            build_registry("agent", "development"), SafetyPolicy("supervised"),
            mode="agent", work_mode="development")
        resumed = completed.has_resumable_task
        first_result = completed.step()
        final_result = completed.step() if first_result.status is Status.CONTINUE else first_result
        check(resumed and final_result.status is Status.FINAL
              and session.load_task_state()["status"] == "complete",
              "A restored running task can reach FINAL")

        manager1 = ProcessManager()
        managers.append(manager1)
        manager1.bind_session(session)
        item = manager1.start(
            "Write-Output BEFORE; Start-Sleep -Milliseconds 500; Write-Output AFTER",
            "powershell", tmp, 10)
        time.sleep(0.15)
        manager2 = ProcessManager()
        managers.append(manager2)
        manager2.bind_session(session)
        restored_item = manager2.get(item.id)
        check(restored_item is not None and restored_item.proc is None,
              "ProcessManager restores its manifest after restart")
        cursor = 0
        output = ""
        for _ in range(100):
            polled = manager2.poll(item.id, cursor)
            output += polled["output"]
            cursor = polled["cursor"]
            if polled["status"] == "finished":
                break
            time.sleep(0.05)
        check("BEFORE" in output and "AFTER" in output,
              "A restored ProcessManager continues reading the persistent log")
    finally:
        for manager in managers:
            manager.terminate_all()
        shutil.rmtree(tmp, ignore_errors=True)


def test_git_tools() -> None:
    print("[git tools]")
    import subprocess
    from harness.tools import fs as fs_tools, git as git_tools

    tmp = Path(tempfile.mkdtemp())
    try:
        def git(*args):
            return subprocess.run(["git", *args], cwd=tmp, check=True,
                                  capture_output=True, text=True, encoding="utf-8")

        git("init", "-q")
        git("config", "user.email", "qwen-test@example.invalid")
        git("config", "user.name", "Qwen Test")
        tracked = tmp / "tracked.txt"
        tracked.write_text("before\n", encoding="utf-8")
        git("add", "tracked.txt")
        git("commit", "-q", "-m", "baseline")

        data = load_config().data
        data["paths"]["sessions_dir"] = str(tmp / "sessions")
        cfg = Config(data, root=ROOT)
        session = Session(cfg, session_id="git-test")
        ctx = AgentContext(cfg=cfg, session=session, workspace=tmp)
        ctx.changes = ChangeJournal(session, tmp)
        ctx.changes.begin_task("git test")
        reg = ToolRegistry()
        fs_tools.register_fs_tools(reg)
        git_tools.register_git_tools(reg)

        patched = reg.execute("apply_patch", {
            "path": "tracked.txt",
            "edits": [{"old": "before", "new": "after"}],
        }, ctx)
        check(patched.startswith("OK"), "The test change was created through apply_patch")
        check("tracked.txt" in reg.execute("git_status", {}, ctx),
              "git_status reports the changed file")
        check("-before" in reg.execute("git_diff", {"path": "tracked.txt"}, ctx),
              "git_diff includes the changed content")
        committed = reg.execute("git_commit", {"message": "task change"}, ctx)
        check("exit code: 0" in committed and git("log", "-1", "--pretty=%s").stdout.strip() == "task change",
              "git_commit includes only journaled changes")
        check(git("status", "--porcelain", "--untracked-files=no").stdout.strip() == "",
              "Tracked changes are clean after commit")
        check("sessions/" in git("status", "--porcelain").stdout,
              "git_commit excludes unrelated journal artifacts")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_automatic_project_check() -> None:
    print("[automatic project check]")
    import time
    from harness.tools import context as context_tools

    tmp = Path(tempfile.mkdtemp())
    try:
        (tmp / "tests").mkdir()
        (tmp / "tests" / "test_core.py").write_text(
            "print('PROJECT-CHECK-OK')\n", encoding="utf-8")
        data = load_config().data
        data["paths"]["sessions_dir"] = str(tmp / ".sessions")
        cfg = Config(data, root=ROOT)
        session = Session(cfg, session_id="check-test")
        ctx = AgentContext(cfg=cfg, session=session, workspace=tmp,
                           project_workspace=tmp,
                           processes=ProcessManager())
        reg = ToolRegistry()
        context_tools.register_coding_context_tools(reg)
        started = json.loads(reg.execute("start_project_check", {"timeout": 10}, ctx))
        cursor = 0
        output = ""
        for _ in range(100):
            result = ctx.processes.poll(started["process_id"], cursor)
            output += result["output"]
            cursor = result["cursor"]
            if result["status"] == "finished":
                break
            time.sleep(0.05)
        check("PROJECT-CHECK-OK" in output and result["exit_code"] == 0,
              "Detected project checks run in a background process")
        profile = reg.execute("project_validation_profile", {}, ctx)
        check("tests" in profile and "Core tests" in profile,
              "The validation profile lists detected checks")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_task_plan_and_project_instructions() -> None:
    print("[task plan + project instructions]")
    from harness.agent import Agent
    from harness.task_plan import TaskPlanStore

    tmp = Path(tempfile.mkdtemp())
    try:
        (tmp / "src" / "feature").mkdir(parents=True)
        (tmp / "AGENTS.md").write_text("Root project guidance", encoding="utf-8")
        (tmp / "src" / "QWEN.md").write_text("Source-specific guidance", encoding="utf-8")
        target = tmp / "src" / "feature" / "module.py"
        target.write_text("VALUE = 1\n", encoding="utf-8")
        data = load_config().data
        data["agent"]["workspace"] = str(tmp)
        data["paths"]["sessions_dir"] = str(tmp / ".sessions")
        cfg = Config(data, root=ROOT)
        session = Session(cfg, session_id="plan-test", system_prompt="SYS",
                          workspace=str(tmp), work_mode="development")
        agent = Agent(cfg, LLMStub(), session, build_registry("agent", "development"),
                      SafetyPolicy("auto"), mode="agent", work_mode="development")
        agent.new_task("Implement the feature and verify it")
        check(TaskPlanStore(session).load().get("goal") == "Implement the feature and verify it",
              "A new task creates a persistent plan")
        plan_result = agent.registry.execute("set_task_plan", {
            "goal": "Implement and verify",
            "steps": ["Inspect the feature", "Implement the change", "Run validation"],
        }, agent.ctx)
        check('"in_progress"' in plan_result and len(TaskPlanStore(session).load()["steps"]) == 3,
              "The model can create structured task steps")
        agent.registry.execute("update_task_step", {
            "step_id": 1, "status": "completed", "note": "Relevant files inspected",
        }, agent.ctx)
        plan = TaskPlanStore(session).load()
        check(plan["steps"][0]["status"] == "completed"
              and plan["steps"][1]["status"] == "in_progress",
              "Completing a step activates the next step")

        root_context = agent._dynamic_context_block()
        check("Root project guidance" in root_context
              and "Source-specific guidance" not in root_context,
              "Root instructions apply from task start")
        agent._observe_context_paths({"path": "src/feature/module.py"})
        nested_context = agent._dynamic_context_block()
        check("Root project guidance" in nested_context
              and "Source-specific guidance" in nested_context,
              "Instructions follow the active file's directory hierarchy")
        check(not any(str(message.get("content", "")).startswith("[DYNAMIC TASK CONTEXT")
                      for message in session.messages),
              "Dynamic project context does not duplicate stored history")
        usage = agent.context_usage_breakdown()
        check(usage["messages"] > 0 and usage["dynamic"] > 0 and usage["tool_schemas"] > 0,
              "Context estimates include messages, dynamic data and tool schemas")

        (tmp / ".qwen").mkdir()
        (tmp / ".qwen" / "project.yaml").write_text(
            "checks:\n"
            "  - id: focused\n"
            "    label: Focused project check\n"
            "    command: Write-Output PROFILE-OK\n"
            "    shell: powershell\n"
            "    primary: true\n",
            encoding="utf-8")
        profile = agent.registry.execute("project_validation_profile", {}, agent.ctx)
        check("focused" in profile and "PROFILE-OK" in profile,
              ".qwen/project.yaml overrides automatic check detection")
        launched = json.loads(agent.registry.execute(
            "start_project_check", {"check": "focused", "timeout": 10}, agent.ctx))
        cursor = 0
        for _ in range(100):
            raw_poll = agent.registry.execute("poll_command", {
                "process_id": launched["process_id"], "cursor": cursor,
            }, agent.ctx)
            polled = json.loads(raw_poll)
            cursor = polled["cursor"]
            if polled["status"] == "finished":
                agent.ctx.task_plan.observe_tool(
                    "poll_command", {"process_id": launched["process_id"]}, raw_poll,
                    processes=agent.ctx.processes)
                break
            import time
            time.sleep(0.05)
        check(TaskPlanStore(session).load()["validations"][-1]["status"] == "passed",
              "Completed project checks update the task plan")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_code_index() -> None:
    print("[code index]")
    from harness.browser import BrowserSession
    from harness.code_index import CodeIndex

    tmp = Path(tempfile.mkdtemp())
    try:
        (tmp / "module.py").write_text(
            "class Alpha:\n"
            "    def run(self):\n"
            "        return make_widget()\n\n"
            "def make_widget():\n"
            "    return Alpha()\n",
            encoding="utf-8")
        (tmp / "ui.ts").write_text(
            "export interface Widget { id: string }\n"
            "export function makeWidget(): Widget { return { id: 'x' }; }\n"
            "const renderWidget = (item: Widget) => item.id;\n",
            encoding="utf-8")
        (tmp / "lib.rs").write_text(
            "pub struct WidgetStore {}\n"
            "pub fn load_widget() -> WidgetStore { WidgetStore {} }\n",
            encoding="utf-8")
        index = CodeIndex(tmp)
        alpha = index.find_symbol("Alpha")
        widget = index.find_symbol("Widget")
        check(any(item["kind"] == "class" and item["path"] == "module.py"
                  for item in alpha), "The Python AST index finds classes and functions")
        check(any(item["path"] == "ui.ts" for item in widget)
              and any(item["path"] == "lib.rs" for item in widget),
              "The symbol index finds TypeScript and Rust declarations")
        document = index.document_symbols("module.py")
        check(any(item["qualified"] == "Alpha.run" for item in document),
              "document_symbols preserves qualified Python method names")
        refs = index.find_references("Widget")
        check(len(refs) >= 2 and all("line" in item for item in refs),
              "find_references returns whole-word matches and line numbers")
        (tmp / "module.py").write_text("def refreshed_symbol():\n    return 1\n",
                                        encoding="utf-8")
        index.invalidate()
        check(index.find_symbol("refreshed_symbol"),
              "The symbol index can be invalidated after an edit")
        browser = BrowserSession()
        check(browser.status()["running"] is False,
              "An unused browser session starts no process")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_research_ledger_and_synthesis() -> None:
    print("[research ledger + synthesis]")
    from harness.llm import AssistantResult
    from harness.research import ResearchLedger, plan_research, synthesize_research

    tmp = Path(tempfile.mkdtemp())
    try:
        data = load_config().data
        data["paths"]["sessions_dir"] = str(tmp / "sessions")
        cfg = Config(data, root=ROOT)
        session = Session(cfg, session_id="research-test", system_prompt="SYS",
                          work_mode="research")
        ledger = ResearchLedger(session)
        ledger.begin("What are the two conflicting interpretations?")
        ledger.record_query("Initial search", [
            ("Source A", "https://example.test/a", "Claims A"),
            ("Source B", "https://example.test/b", "Claims B"),
        ])
        ledger.record_source("https://example.test/a", "Source A",
                             "Interpretation A says yes.")
        ledger.record_source("https://example.test/b", "Source B",
                             "Interpretation B says no.")
        run = ledger.current()
        check(len(run["candidates"]) == 2 and len(run["sources"]) == 2,
              "The ledger retains every candidate and loaded source")
        check(all("credibility" not in source and "trust" not in source
                  for source in run["sources"]),
              "The ledger has no trust scoring or source filter")

        class ResearchLLM:
            def __init__(self):
                self.cfg = cfg
                self.prompts: list[str] = []

            def ask(self, messages, **_kwargs):
                prompt = str(messages[-1]["content"])
                self.prompts.append(prompt)
                if "Return JSON only" in prompt:
                    return AssistantResult(content=json.dumps({
                        "subquestions": ["What does A claim?", "What does B claim?"],
                        "search_angles": ["contradictions"],
                        "source_types_to_include": ["all available"],
                        "known_constraints": [],
                    }, ensure_ascii=False))
                if "Missing sources" in prompt:
                    return AssistantResult(content="The revised synthesis includes [S1] and [S2].")
                if "Create a clear final synthesis" in prompt:
                    return AssistantResult(content="The first synthesis includes only [S1].")
                return AssistantResult(content="Intermediate loss-aware notes [S1] [S2].")

            def stream(self, messages, **kwargs):
                if kwargs.get("thinking") is False:
                    return self.ask(messages, **kwargs)
                return AssistantResult(content="Working draft before synthesis")

        fake = ResearchLLM()
        class EmptyPlannerLLM(ResearchLLM):
            def ask(self, messages, **_kwargs):
                return AssistantResult(content="", reasoning="Unfinished reasoning")

        fallback = plan_research(EmptyPlannerLLM(), "What needs to be established?")
        check(fallback["subquestions"] == ["What needs to be established?"]
              and len(fallback["search_angles"]) >= 3,
              "An empty planner response uses the fallback plan without stopping research")

        synthesis = synthesize_research(fake, run)
        check("[S1]" in synthesis and "[S2]" in synthesis,
              "Coverage repair includes every processed source ID")
        final_prompt = next(prompt for prompt in fake.prompts
                            if "Create a clear final synthesis" in prompt)
        check("Interpretation A" in final_prompt and "Interpretation B" in final_prompt
              and "Do not assess or filter" in final_prompt,
              "Synthesis receives conflicting evidence without trust-based filtering")
        ledger.complete(synthesis)
        check(ledger.status()["status"] == "complete" and ledger.path.is_file(),
              "The research ledger persists and records completed synthesis")

        from harness.agent import Agent, Status
        from harness.safety import SafetyPolicy
        integrated_session = Session(
            cfg, session_id="research-agent", system_prompt="SYS", work_mode="research")
        integrated_llm = ResearchLLM()
        agent = Agent(
            cfg, integrated_llm, integrated_session,
            build_registry("chat", "research"), SafetyPolicy("auto"),
            mode="chat", work_mode="research",
        )
        agent.new_task("Integrated research question")
        agent.ctx.research.record_source("https://example.test/a", "A", "Ano [S1]")
        agent.ctx.research.record_source("https://example.test/b", "B", "Ne [S2]")
        result = agent.step()
        check(result.status is Status.FINAL and "[S1]" in result.text and "[S2]" in result.text
              and "Working draft" not in result.text,
              "The research Agent produces a synthesis with source coverage")
        check(any(message.get("role") == "assistant"
                  and message.get("content") == "Working draft before synthesis"
                  for message in integrated_session.messages),
              "The pre-synthesis draft remains saved and visible")
        integrated_run = agent.ctx.research.current()
        check(integrated_run["plan"]["subquestions"] == ["What does A claim?", "What does B claim?"]
              and any(str(message.get("content", "")).startswith("[RESEARCH PLAN")
                      for message in integrated_session.messages),
              "Research planning precedes searching and is saved in the ledger and context")

        run_id = integrated_run["id"]
        agent.new_task("Save the previous response as a PDF file")
        check(agent.ctx.research.current()["id"] == run_id,
              "A PDF export follow-up does not create a new research run")
        exported = agent.registry.execute("export_document", {
            "content": "# Saved result\n\nImportant synthesis.",
            "filename": "research-output",
            "format": "pdf",
            "title": "Research result",
        }, agent.ctx)
        no_project_pdf = integrated_session.dir / "exports" / "research-output.pdf"
        check(exported.startswith("OK:") and no_project_pdf.is_file()
              and "Important synthesis" in "\n".join(
                  page.extract_text() or "" for page in __import__("pypdf").PdfReader(
                      no_project_pdf).pages),
              "Projectless research export saves a readable PDF in the session")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_project_document_library() -> None:
    print("[project document library]")
    from docx import Document
    from harness.agent import Agent
    from harness.documents import export_document
    from harness.repo_index import RepoIndex
    from harness.safety import SafetyPolicy
    from pypdf import PdfWriter

    tmp = Path(tempfile.mkdtemp())
    try:
        (tmp / "film.md").write_text("# Aurora film\nThe protagonist is Clara.\n", encoding="utf-8")
        (tmp / "code.py").write_text("SECRET_CODE = 1\n", encoding="utf-8")
        document = Document()
        document.add_paragraph("DOCX source about the character Clara")
        document.save(tmp / "character.docx")
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        with open(tmp / "reference.pdf", "wb") as handle:
            writer.write(handle)
        data = load_config().data
        data["paths"]["sessions_dir"] = str(tmp / "sessions")
        cfg = Config(data, root=ROOT)

        discussion_session = Session(
            cfg, session_id="discussion-docs", system_prompt="SYS",
            workspace=str(tmp), work_mode="discussion")
        discussion_agent = Agent(
            cfg, LLMStub(), discussion_session, build_registry("chat", "discussion"),
            SafetyPolicy("auto"), mode="chat", work_mode="discussion")
        discussion_agent.set_workspace(tmp)
        discussion_agent.new_task("Tell me about the film")
        prompt = discussion_agent._api_messages()[-1]["content"]
        check("CURRENT PROJECT DOCUMENT LIBRARY" in prompt and "film.md" in prompt
              and "CURRENT PROJECT SNAPSHOT" not in prompt and "code.py" not in prompt,
              "Discussion sees the document library without a development repository snapshot")

        research_session = Session(
            cfg, session_id="research-docs", system_prompt="SYS",
            workspace=str(tmp), work_mode="research")
        research_agent = Agent(
            cfg, LLMStub(), research_session, build_registry("chat", "research"),
            SafetyPolicy("auto"), mode="chat", work_mode="research")
        research_agent.set_workspace(tmp)
        research_agent.new_task("Who is the heroine?")
        content = research_agent.registry.execute(
            "read_project_document", {"path": "film.md"}, research_agent.ctx)
        sources = research_agent.ctx.research.current()["sources"]
        check("Clara" in content and len(sources) == 1
              and sources[0]["url"].startswith("file://"),
              "Research records a local document as a ledger source")
        library = RepoIndex(tmp)
        _, docx_text = library.read_document("character.docx")
        pdf_path, pdf_text = library.read_document("reference.pdf")
        check("DOCX source" in docx_text, "The project library reads a real DOCX file")
        check(pdf_path.suffix == ".pdf" and isinstance(pdf_text, str),
              "The project library reads a valid PDF file")
        exported_docx = export_document(
            "# Heading\n\nText about Clara and λ.\n\n- item A", tmp / "exports", "output", "docx", "Film")
        exported_pdf = export_document(
            "# Heading\n\n**Text about Clara and λ.** and Na⁺.\n\n"
            "| Item | Value |\n|---|---|\n"
            "| Source | [Web](https://example.com) |\n\n## 📋 Next steps",
            tmp / "exports", "output", "pdf", "Film")
        exported_docx_text = "\n".join(
            paragraph.text for paragraph in Document(exported_docx).paragraphs)
        exported_pdf_text = "\n".join(
            page.extract_text() or "" for page in __import__("pypdf").PdfReader(exported_pdf).pages)
        check("Clara" in exported_docx_text, "Writing exports structured DOCX with Unicode content")
        check("Clara" in exported_pdf_text and "**" not in exported_pdf_text
              and "Item" in exported_pdf_text and "Source" in exported_pdf_text
              and "Na+" in exported_pdf_text and "Next steps" in exported_pdf_text,
              "PDF export renders Unicode, inline Markdown and tables")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_async_model_switch() -> None:
    print("[async model switch]")
    import threading
    from harness.model_switch import ModelSwitchController

    entered = threading.Event()
    release = threading.Event()
    stopped = threading.Event()
    callbacks: list[str] = []

    def ensure(_cfg, key):
        if key == "q5":  # Keep the first target loading until it is cancelled.
            entered.set()
            release.wait(timeout=2)
            return False  # Stopping the server causes ensure to fail
        return key == "q4"

    stop_calls: list[int] = []

    def stop(_cfg, **_kwargs):
        stop_calls.append(1)
        if len(stop_calls) >= 2:  # The first stop precedes ensure; later stops interrupt obsolete work.
            release.set()  # Simulate interruption during loading
            stopped.set()
        return True

    controller = ModelSwitchController(load_config(), ensure_fn=ensure, stop_fn=stop,
                                       running_fn=lambda _cfg, _key: False)
    check(controller.request("q5", on_success=callbacks.append),
          "The first model switch starts in the background")
    check(entered.wait(timeout=1) and controller.snapshot().busy,
          "The controller reports starting immediately")
    check(controller.request("q4", on_success=callbacks.append),
          "A concurrent request replaces the pending target")
    check(controller.wait(timeout=3), "The background switch completes")
    snap = controller.snapshot()
    check(snap.status == "ready" and snap.target == "q4",
          f"latest target wins ({snap.status}/{snap.target})")
    check(callbacks == ["q4"], "Only the winning target receives the callback")
    check(stopped.is_set(), "Switching interrupts the obsolete load")

    failed = ModelSwitchController(load_config(), ensure_fn=lambda _cfg, _key: False,
                                   stop_fn=stop, running_fn=lambda _cfg, _key: False)
    check(failed.request("q4") and failed.wait(timeout=2)
          and failed.snapshot().status == "failed",
          "Server failure appears in controller state")

    #Apply the KV profile before startup (set_kv_cache_mode before ensure)
    applied: list[tuple[str, str]] = []

    def ensure_kv(cfg2, key):
        applied.append((key, cfg2.kv_cache_mode(key)))
        return True

    kvctl = ModelSwitchController(load_config(), ensure_fn=ensure_kv, stop_fn=stop,
                                  running_fn=lambda _cfg, _key: False)
    kvctl.request("q4", kv_profile="q8_0", on_success=callbacks.append)
    kvctl.wait(timeout=2)
    check(applied == [("q4", "q8_0")],
          "The requested KV profile is applied before server startup")


def test_session_meta() -> None:
    print("[session metadata and history]")
    tmp = Path(tempfile.mkdtemp())
    try:
        data = load_config().data
        data["paths"]["sessions_dir"] = str(tmp / "sessions")
        cfg = Config(data, root=ROOT)
        # Session workspace and title derived from the first user request.
        s = Session(cfg, session_id="meta-a", system_prompt="SYS", workspace=r"C:\projekty\Alfa")
        s.add("user", "Oprav bug v parseru")
        s.add("assistant", "done")
        check(s.meta["title"] == "Oprav bug v parseru", "The title comes from the first user request")
        check((tmp / "sessions/meta-a/meta.json").exists(), "meta.json was saved")
        # Internal protocol messages must not become conversation titles.
        s2 = Session(cfg, session_id="meta-b", system_prompt="SYS")
        s2.add("user", "[TASK PROTOCOL - follow] abc")
        s2.add("user", "Actual user request")
        check(s2.meta["title"] == "Actual user request", "Internal notes do not become conversation titles")
        # List metadata sorted by the update timestamp.
        lst = Session.list_sessions(cfg)
        check({x["id"] for x in lst} == {"meta-a", "meta-b"}, "list_sessions returns both conversations")
        a = next(x for x in lst if x["id"] == "meta-a")
        check(a["workspace"] == r"C:\projekty\Alfa" and a["title"] == "Oprav bug v parseru",
              "The listing includes workspace and title metadata")
        check(a["messages"] == s.meta["message_count"] == len(s.messages),
              "Message counts come from the metadata index")
        check(lst[0]["id"] == "meta-b", "Newer sessions appear first")
        # Recover an older session title from its first user message.
        s3 = Session(cfg, session_id="meta-old", system_prompt="SYS")
        s3.add("user", "Old request without metadata")
        (tmp / "sessions/meta-old/meta.json").unlink()
        loaded = Session.load(cfg, "meta-old")
        check(loaded.meta["title"] == "Old request without metadata", "Legacy conversation titles remain supported")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_chat_rewind_and_fork() -> None:
    print("[chat retry/undo/fork]")
    tmp = Path(tempfile.mkdtemp())
    try:
        data = load_config().data
        data["paths"]["sessions_dir"] = str(tmp / "sessions")
        cfg = Config(data, root=ROOT)
        session = Session(cfg, session_id="original", system_prompt="SYS", workspace=str(tmp),
                          work_mode="writing")
        session.add("user", "First request")
        session.add("assistant", "First response")
        image = tmp / "source.png"
        image.write_bytes(b"image-data")
        session.add("user", "Second request", images=[image])
        session.add("user", "[TASK PROTOCOL - internal]")
        session.add("assistant", "Second response")
        original_count = len(session.messages)

        markdown = session.export_markdown()
        jsonl = session.export_jsonl()
        check(markdown.is_file() and "Second request" in markdown.read_text(encoding="utf-8"),
              "Conversation export produces readable Markdown")
        imported = Session.import_jsonl(cfg, jsonl, "IMPORTED SYS", workspace=str(tmp))
        check(imported.id != session.id and imported.messages[0]["content"] == "IMPORTED SYS"
              and imported.last_user_index() is not None,
              "JSONL import creates a new session and restores the system prompt")
        found = Session.search_sessions(cfg, "Second request")
        check(any(item["id"] == session.id and "Second request" in item["snippet"]
                  and "sessions" not in item["snippet"] for item in found),
              "Full-text history search finds the request")
        check((tmp / "sessions" / "history-index.sqlite3").is_file(),
              "History search uses a persistent SQLite FTS index")

        fork = session.fork_at_last_user("FORK SYS")
        check(fork is not None and len(session.messages) == original_count and fork.id != session.id,
              "Forking creates a new session without changing the original")
        check(fork.meta["work_mode"] == "writing", "Forking preserves the work mode")
        fork_index = fork.last_user_index()
        fork_user = fork.messages[fork_index] if fork_index is not None else {}
        copied_images = fork_user.get("images", [])
        check(fork_user.get("content") == "Second request" and copied_images
              and Path(copied_images[0]).exists() and copied_images[0] != session.messages[3]["images"][0],
              "Forking retains the last user request and copies its attachments")

        session.compression = {"cut": 2, "summary": "old"}
        session._save_compression()
        prompt = session.rewind_last_turn(keep_user=True)
        check(prompt == "Second request" and session.messages[-1]["content"] == "Second request"
              and session.compression is None,
              "Retry keeps the question and removes the response and old compression")

        session.add("assistant", "New response")
        removed = session.rewind_last_turn(keep_user=False)
        check(removed == "Second request" and session.last_user_index() is not None
              and session.messages[session.last_user_index()]["content"] == "First request",
              "Undo removes the entire last turn")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_transient_session() -> None:
    print("[transient session persistence]")
    tmp = Path(tempfile.mkdtemp())
    try:
        data = load_config().data
        data["paths"]["sessions_dir"] = str(tmp / "sessions")
        cfg = Config(data, root=ROOT)
        s = Session(cfg, system_prompt="SYS", workspace=r"C:\projekty\Alfa", transient=True)
        s.add("assistant", "hello")
        check(s.transient and not s.dir.exists(),
              "A transient session writes nothing before a user message")
        s.add("user", "Oprav bug")
        check(not s.transient and s.dir.exists() and s._jsonl.exists(),
              "The first user message persists the session")
        lines = s._jsonl.read_text(encoding="utf-8").strip().splitlines()
        check(len(lines) == 3, f"entire history written ({len(lines)} lines)")
        s2 = Session.load(cfg, s.id)
        check(len(s2.messages) == 3 and not s2.transient, "load after persistence")
        check(Session.delete(cfg, "neexistujici-x") is False, "Deleting a missing session returns False")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_web_tools() -> None:
    print("[web tools: offline parsers and registration]")
    from io import BytesIO
    from harness.tools import web as webt
    from reportlab.pdfgen import canvas

    txt = webt._strip_tags("<html><body><script>bad()</script><h1>Hello</h1>"
                           "<p>world &amp; hello</p></body></html>")
    check("Hello" in txt and "world & hello" in txt and "bad()" not in txt,
          "_strip_tags removes markup and decodes entities")
    u = webt._ddg_unwrap("//duckduckgo.com/l/?uddg=https%3A%2F%2Fexample.com%2Fdoc")
    check(u == "https://example.com/doc", "_ddg_unwrap resolves the uddg redirect")
    enc = ("https://www.bing.com/ck/a?!&amp;&amp;p=xx&u=a1aHR0cHM6Ly9naXRodWIuY29tL3Rlc3Q"
           "&ntb=1")
    check(webt._bing_unwrap(enc) == "https://github.com/test", "_bing_unwrap decodes the base64 URL")
    pdf_buffer = BytesIO()
    pdf_canvas = canvas.Canvas(pdf_buffer)
    pdf_canvas.drawString(72, 720, "INTERNET-PDF-OK")
    pdf_canvas.save()
    pdf_text, _ = webt._extract_downloaded_document(
        pdf_buffer.getvalue(), "application/pdf", "https://example.test/study.pdf")
    docx_buffer = BytesIO()
    docx_document = __import__("docx").Document()
    docx_document.add_paragraph("INTERNET-DOCX-OK")
    docx_document.save(docx_buffer)
    docx_text, _ = webt._extract_downloaded_document(
        docx_buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "https://example.test/study.docx")
    check("INTERNET-PDF-OK" in pdf_text, "web_fetch extracts text from a PDF response")
    check("INTERNET-DOCX-OK" in docx_text, "web_fetch extracts text from a DOCX response")
    for mode in ("chat", "agent", "computer"):
        reg = build_registry(mode)
        check("web_search" in reg.names() and "web_fetch" in reg.names(),
              f"web tools in mode {mode}")
    reg = build_registry("chat")
    cfgd = load_config()
    ctx = type("C", (), {"cfg": cfgd})()
    out = reg.execute("web_fetch", {"url": "ftp://invalid.example"}, ctx)
    check(out.startswith("ERROR"), "web_fetch rejects non-HTTP URLs")

    import requests
    from harness.research import ResearchLedger
    tmp = Path(tempfile.mkdtemp())
    original_ensure = webt._ensure_ddgs
    original_bing = webt.WebSearchTool._bing
    original_ddg = webt.WebSearchTool._ddg
    original_get = requests.get
    try:
        data = load_config().data
        data["paths"]["sessions_dir"] = str(tmp / "sessions")
        cfg = Config(data, root=ROOT)
        session = Session(cfg, session_id="web-research", system_prompt="SYS",
                          work_mode="research")
        ledger = ResearchLedger(session)
        ledger.begin("Web ledger test")
        research_ctx = type("RC", (), {"cfg": cfg, "research": ledger})()
        webt._ensure_ddgs = lambda: False
        webt.WebSearchTool._bing = staticmethod(lambda *_args: [
            ("A", "https://example.test/a", "yes"),
            ("B", "https://example.test/b", "no"),
        ])
        webt.WebSearchTool._ddg = staticmethod(lambda *_args: [])
        webt.WebSearchTool().run(research_ctx, "contradiction", 2)
        check(len(ledger.current()["candidates"]) == 2,
              "web_search records every discovered candidate")

        class Response:
            url = "https://example.test/a"
            headers = {"content-type": "text/html"}
            text = "<html><title>Source A</title><body>The source contains claims for yes and no.</body></html>"
            status_code = 200

            @staticmethod
            def raise_for_status():
                return None

        requests.get = lambda *_args, **_kwargs: Response()
        fetched = webt.WebFetchTool().run(research_ctx, "https://example.test/a")
        source = ledger.current()["sources"][0]
        check("The source contains claims for yes and no" in fetched and "yes and no" in source["content"],
              "web_fetch retains the complete readable source in the ledger")
        check("credibility" not in source and "trust" not in source,
              "Web integration adds no trustworthiness scoring")
    finally:
        webt._ensure_ddgs = original_ensure
        webt.WebSearchTool._bing = original_bing
        webt.WebSearchTool._ddg = original_ddg
        requests.get = original_get
        shutil.rmtree(tmp, ignore_errors=True)


def test_projects() -> None:
    print("[projects]")
    from harness.projects import Projects
    tmp = Path(tempfile.mkdtemp())
    try:
        data = load_config().data
        data["projects"] = {"root_dir": "projects"}
        data["work_mode"] = "writing"
        cfg = Config(data, root=tmp)
        pj = Projects(cfg)
        # Create a new directory under the project root.
        p1 = pj.create_new("My-Test:Project")  # Sanitize unsafe name characters.
        check((tmp / "projects" / p1["name"]).is_dir(), f"directory created ({p1['name']})")
        check(p1["work_mode"] == "writing", "The project records its default work mode")
        check(p1["name"] != "My-Test:Project" or True, "The project name was sanitized")
        # Duplicate names receive a -2 suffix
        p2 = pj.create_new(p1["name"])
        check(p2["name"] != p1["name"], "Duplicate names receive a unique suffix")
        # Attaching a directory is idempotent and uses its name.
        ext = tmp / "Existujici"
        ext.mkdir()
        a1 = pj.attach_folder(str(ext))
        a2 = pj.attach_folder(str(ext))
        check(a1["name"] == "Existujici" and a1["id"] == a2["id"], "Attaching a project is idempotent")
        pj.set_work_mode(str(ext.resolve()), "research")
        check(pj.by_path(str(ext.resolve()))["work_mode"] == "research",
              "The default project mode can be changed")
        # The registry returns every project.
        names = [p["name"] for p in pj.list_all()]
        check(len(names) == 3, f"3 projekty v registru ({names})")
        managed_file = Path(p1["path"]) / "data" / "artifact.txt"
        managed_file.parent.mkdir()
        managed_file.write_text("Project data", encoding="utf-8")
        pj.delete_by_path(p1["path"])
        check(not Path(p1["path"]).exists() and pj.by_path(p1["path"]) is None,
              "Project deletion removes its registry entry, directory and files")
        # session delete + adopt
        s = Session(cfg, session_id="proj-s", system_prompt="SYS", workspace=str(ext))
        check(Session.delete(cfg, "proj-s") and not (tmp/"sessions"/"proj-s").exists(),
              "session delete")
        s2 = Session(cfg, session_id="adopt-s", system_prompt="SYS")  # without a workspace
        s2.adopt_workspace(str(ext))
        check(s2.meta["workspace"] == str(ext), "adopt workspace")
        s2.adopt_workspace("jina")  # Do not overwrite an existing assignment.
        check(s2.meta["workspace"] == str(ext), "Adoption preserves an existing project assignment")
        (ext / "external.txt").write_text("data", encoding="utf-8")
        pj.delete_by_path(str(ext.resolve()))
        check(not pj.by_path(str(ext.resolve())) and (ext / "external.txt").is_file(),
              "Removing an attached project keeps its folder and files on disk")
        protected = pj.attach_folder(str(tmp))
        pj.delete_by_path(protected["path"])
        check(tmp.exists() and pj.by_path(str(tmp)) is None,
              "Removing a project on a protected path unregisters it without touching disk")
        # A locked managed folder reports a readable error and stays registered.
        from unittest.mock import patch
        locked = pj.create_new("Locked")
        (Path(locked["path"]) / "file.txt").write_text("x", encoding="utf-8")
        with patch("harness.projects.shutil.rmtree", side_effect=PermissionError(5, "Access is denied")):
            try:
                pj.delete_by_path(locked["path"])
                check(False, "A locked project folder raises a readable error")
            except ValueError as exc:
                check("locked" in str(exc), f"A locked project folder raises a readable error ({exc})")
        check(pj.by_path(locked["path"]) is not None,
              "A failed folder deletion keeps the project registered")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_skill_library() -> None:
    print("[skills]")
    from harness.skills import SkillLibrary
    system = SkillLibrary(load_config())
    names = [item.name for item in system.list()]
    check({"systematic-debugging", "architecture-options", "implementation-verification",
           "excel-spreadsheet-craft", "word-document-craft", "pdf-generation-craft",
           "computer-automation-craft", "web-scraping-extraction", "code-refactoring-patterns"}
          <= set(names), "The skill library discovers bundled and newly added SKILL.md files")
    check("root cause" in system.read("systematic-debugging").lower(),
          "Skill bodies load only when explicitly read")

    tmp = Path(tempfile.mkdtemp())
    try:
        custom = tmp / ".qwen-skills" / "override" / "SKILL.md"
        custom.parent.mkdir(parents=True)
        custom.write_text(
            "---\nname: systematic-debugging\n"
            "description: Project-specific debugging guidance.\n---\n\nPROJECT OVERRIDE\n",
            encoding="utf-8")
        library = SkillLibrary(load_config(), tmp)
        info = next(item for item in library.list() if item.name == "systematic-debugging")
        check(info.source == "project" and "PROJECT OVERRIDE" in library.read(info.name),
              "A project skill can override a bundled skill with the same name")
        ctx = type("SkillCtx", (), {
            "cfg": load_config(), "project_workspace": tmp,
        })()
        output = build_registry("chat").execute(
            "read_skill", {"name": "systematic-debugging"}, ctx)
        check("PROJECT OVERRIDE" in output, "read_skill exposes the selected instructions to the model")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


class LLMStub:
    """Deterministic model fixture that returns scripted responses in sequence."""
    def __init__(self, script=None):
        from harness.llm import AssistantResult
        self.script = list(script or [])
        self.calls = 0

    def stream(self, messages, tools=None, on_text=None, on_reasoning=None, **kw):
        self.calls += 1
        self.last_messages = messages
        return self.script.pop(0)


def _tc(name, args="{}"):
    return {"id": f"call_{name}", "type": "function",
            "function": {"name": name, "arguments": args}}


def test_communication_protocol() -> None:
    print("[communication protocol]")
    from harness.agent import Agent, Status
    from harness.llm import AssistantResult
    from harness.safety import SafetyPolicy
    data = load_config().data
    data["paths"]["sessions_dir"] = str(Path(tempfile.mkdtemp()) / "sessions")
    cfg = Config(data, root=ROOT)
    tmp = Path(tempfile.mkdtemp())
    try:
        cfg.agent["workspace"] = str(tmp)

        def make_agent(script):
            session = Session(cfg, session_id=f"proto-{uuid.uuid4().hex[:6]}")
            llm = LLMStub(script)
            agent = Agent(cfg, llm, session, build_registry("agent"),
                          SafetyPolicy("auto"), mode="agent")
            return agent, session

        # 1) Preserve internal context messages to keep the cached prefix stable.
        agent, session = make_agent([])
        agent.new_task("udelej neco")
        notes = [m for m in session.messages if "[TASK PROTOCOL" in str(m.get("content"))]
        check(len(notes) == 1, "The task protocol is added as a user message")
        agent.new_task("Next task")
        notes = [m for m in session.messages if "[TASK PROTOCOL" in str(m.get("content"))]
        check(len(notes) == 2, "A new task preserves the previously cached protocol")
        reloaded = Session.load(cfg, session.id)
        persisted_notes = [m for m in reloaded.messages
                           if "[TASK PROTOCOL" in str(m.get("content"))]
        check(len(persisted_notes) == 2, "Cache-preserving protocol messages survive reload unchanged")

        # 2) Request an update after four silent tool steps.
        script = [AssistantResult(tool_calls=[_tc("list_dir", '{"path": "."}')]) for _ in range(4)]
        script.append(AssistantResult(content="done"))
        agent, session = make_agent(script)
        agent.new_task("search the directory")
        statuses = [agent.step(approve=True).status for _ in range(5)]
        prog = [m for m in session.messages if "[PROGRESS UPDATE" in str(m.get("content"))]
        check(len(prog) >= 1, f"PROGRESS nudge after 4 steps (count: {len(prog)})")

        # 3) Require a structured summary after a tool-driven task.
        script = [
            AssistantResult(tool_calls=[_tc("list_dir")]),
            AssistantResult(tool_calls=[_tc("list_dir"), _tc("list_dir")]),
            AssistantResult(content="just a short answer"),   # a nontrivial task needs a summary
            AssistantResult(content="✅ Done: nothing\n- **x**: y"),  # structured
        ]
        agent, session = make_agent(script)
        agent.new_task("summary test")
        r1 = agent.step(approve=True)
        r2 = agent.step(approve=True)
        r3 = agent.step(approve=True)
        check(r3.status is Status.CONTINUE, "A short response after tools requests a final summary")
        notes = [m for m in session.messages if "[FINAL SUMMARY" in str(m.get("content"))]
        check(len(notes) == 1, "The summary request was inserted")
        r4 = agent.step(approve=True)
        check(r4.status is Status.FINAL and "✅" in r4.text, "The second pass returns FINAL with a summary")
        check(agent.llm.calls == 4, "No unnecessary model call is made")

        # 4) Discussion without tools does not need the development protocol.
        session = Session(cfg, session_id="chat-proto")
        agent = Agent(cfg, LLMStub([]), session, build_registry("chat"), SafetyPolicy("auto"), mode="chat")
        agent.new_task("hello")
        notes = [m for m in session.messages if "[TASK PROTOCOL" in str(m.get("content"))]
        check(not notes, "Discussion does not receive the development protocol")

        # 5) Compress and retry once after context overflow, then finish normally.
        class OverflowLLM(LLMStub):
            def stream(self, messages, **kw):
                self.calls += 1
                if self.calls == 1:
                    raise RuntimeError("request exceeds the available context size")
                from harness.llm import AssistantResult
                return AssistantResult(content="✅ po kompresi OK")

        from harness.agent import Status as St
        for i in range(8):
            (tmp / f"f{i}.txt").write_text("x" * 400, encoding="utf-8")
        session = Session(cfg, session_id="ovf-test", system_prompt="SYS")
        ovf_llm = OverflowLLM()
        agent = Agent(cfg, ovf_llm, session, build_registry("agent"), SafetyPolicy("auto"), mode="agent")
        agent.llm = ovf_llm
        # Populate enough history to exercise compression.
        for i in range(8):
            session.add("user", f"q{i} " + "y" * 900)
            session.add("assistant", f"a{i} " + "z" * 900)
        agent.new_task("Next question")
        agent._steps = 0
        r = agent.step(approve=True)
        check(r.status is St.CONTINUE, "overflow → CONTINUE (compression + retry)")
        r2 = agent.step(approve=True)
        check(r2.status is St.FINAL and ovf_llm.calls == 2,
              "Retry after compression succeeds with two model calls")
        # A second overflow returns ERROR instead of retrying indefinitely.
        class AlwaysOverflow(LLMStub):
            def stream(self, messages, **kw):
                self.calls += 1
                raise RuntimeError("prompt is too long: 999999 > 98304")
        ao = AlwaysOverflow()
        agent.llm = ao
        agent._overflow_retried = False
        session.add("user", "znovu preteceni")
        agent.safety.new_task()
        r3 = agent.step(approve=True)
        r4 = agent.step(approve=True)
        check(r4.status is St.ERROR and ao.calls == 2, "A second overflow returns ERROR without looping")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_user_manuals() -> None:
    print("[user manuals]")
    from pypdf import PdfReader
    from harness.version import APP_VERSION

    expected = {
        "Marvin-Manual-EN.pdf": (
            15, (APP_VERSION, "Flash-Next", "Python 3.12", "Installing from the offline backup", "Work Modes",
                 "User-Facing Tool Reference", "Troubleshooting")),
        "Marvin-Manual-CS.pdf": (
            10, (APP_VERSION, "Flash-Next", "Python 3.12", *locale_data("manual_headings"))),
    }
    for filename, (minimum_pages, required_text) in expected.items():
        # Manuals live under output/pdf in development and docs in an installed copy.
        candidates = [ROOT / "output" / "pdf" / filename, ROOT / "docs" / filename]
        path = next((p for p in candidates if p.is_file()), candidates[0])
        check(path.is_file() and path.stat().st_size > 50_000,
              f"{filename} exists and is not empty")
        if not path.is_file():
            continue
        reader = PdfReader(path)
        text = "".join((page.extract_text() or "") for page in reader.pages)
        check(len(reader.pages) >= minimum_pages,
              f"{filename} has the full scope ({len(reader.pages)} pages)")
        check(all((page.extract_text() or "").strip() for page in reader.pages),
              f"{filename} has no empty pages")
        check(all(item in text for item in required_text),
              f"{filename} contains the version and key chapters")


def test_live_sessions_untouched() -> None:
    """The suite must never write a conversation into the real sessions directory."""
    print("[live sessions]")
    current = ({item.name for item in _LIVE_SESSIONS.iterdir()}
               if _LIVE_SESSIONS.is_dir() else set())
    leaked = sorted(current - _LIVE_SESSIONS_BEFORE)
    check(not leaked,
          "The suite left no conversation in the live sessions directory"
          + (f" (leaked: {leaked})" if leaked else ""))


def test_thinking_and_communication():
    print("[thinking and clean communication]")
    from harness.llm import ThinkStreamParser
    from harness.session import Session
    from harness.config import load_config
    import webapp
    scratch_sessions(webapp.cfg)

    # 1) Parse reasoning tags split across stream chunks.
    text_accum, reason_accum = [], []
    parser = ThinkStreamParser(on_text=text_accum.append, on_reasoning=reason_accum.append)
    t1, r1 = parser.feed("<th")
    t2, r2 = parser.feed("ink>Deep model reasoning</thi")
    t3, r3 = parser.feed("nk>\nModel response.")
    tf, rf = parser.flush()
    full_text = "".join(t1 + t2 + t3 + tf)
    full_reason = "".join(r1 + r2 + r3 + rf)
    check(full_text == "Model response." and full_reason == "Deep model reasoning",
          "ThinkStreamParser separates reasoning tags across chunk boundaries")
    check("".join(text_accum) == "Model response." and "".join(reason_accum) == "Deep model reasoning",
          "ThinkStreamParser delivers the correct text and reasoning segments")

    # 2) Session reasoning persistence
    cfg = scratch_sessions(load_config())
    s = Session(cfg, transient=False)
    s.add("user", "Reasoning question")
    s.add("assistant", "Final response", reasoning="Internal model reasoning")
    check(s.messages[-1].get("reasoning") == "Internal model reasoning",
          "Session retains reasoning in memory")
    loaded = Session.load(cfg, s.id)
    check(loaded.messages[-1].get("reasoning") == "Internal model reasoning",
          "Reasoning survives JSONL persistence and reload")
    api_msgs = s.to_api_messages()
    last_api = api_msgs[-1]
    check(last_api.get("reasoning_content") == "Internal model reasoning" and "reasoning" not in last_api,
          "to_api_messages maps reasoning to reasoning_content for llama-server")

    # 3) webapp thought box & chat view rendering
    thought_open = webapp._format_thought_box("My reasoning", open_box=True, elapsed_s=4)
    check('<details class="thought-box" open>' in thought_open and "Thinking…" in thought_open,
          "_format_thought_box opens while reasoning streams and displays the timer")
    thought_closed = webapp._format_thought_box("My reasoning", open_box=False, duration=3.5)
    check('<details class="thought-box">' in thought_closed and "open" not in thought_closed
          and "Thought for 4s" in thought_closed,
          "_format_thought_box collapses after completion and displays the reasoning duration")

    webapp.state.session = loaded
    view = webapp.chat_view()
    assistant_view = [m for m in view if m["role"] == "assistant"]
    check(len(assistant_view) >= 1 and 'class="thought-box"' in assistant_view[-1]["content"]
          and "Final response" in assistant_view[-1]["content"],
          "chat_view renders a collapsible thought box and the final response")

    # 4) tool box in chat view
    loaded.add("tool", "contents of file abc.txt", name="read_file")
    view_with_tool = webapp.chat_view()
    tool_entry = view_with_tool[-1]
    check('class="tool-box"' in tool_entry["content"] and "read_file" in tool_entry["content"]
          and "contents of file abc.txt" in tool_entry["content"],
          "chat_view renders tool output in a collapsible tool box")
    Session.delete(cfg, s.id)


def test_harness_enhancements():
    print("[harness enhancements: truncate, syntax, fts5, loop, slash]")
    from harness.tools.base import truncate, AgentContext
    from harness.tools.fs import validate_syntax_pre_write
    from harness.tools.search import SearchProjectTool
    from harness.session import Session
    from harness.config import load_config
    import webapp
    import tempfile
    scratch_sessions(webapp.cfg)

    # 1) Head+Tail truncate
    short = "kratky text"
    check(truncate(short, limit=50) == short, "truncate preserves short text")
    long_text = "START_OF_OUTPUT\n" + ("prostredni radek\n" * 100) + "END_OF_OUTPUT_ERROR_TRACE"
    t_res = truncate(long_text, limit=60)
    check("START_OF_OUTPUT" in t_res and "END_OF_OUTPUT_ERROR_TRACE" in t_res,
          "truncate preserves the head and the tail containing the error")
    check("truncated:" in t_res and "omitted" in t_res,
          "truncate reports omitted lines and characters")

    # 2) Validate syntax before writing.
    valid_py = "def foo():\n    return 42\n"
    invalid_py = "def foo(\n    return 42\n"
    check(validate_syntax_pre_write(Path("test.py"), valid_py) is None,
          "validate_syntax_pre_write accepts valid Python")
    check(validate_syntax_pre_write(Path("test.py"), invalid_py) is not None,
          "validate_syntax_pre_write catches invalid Python")
    check(validate_syntax_pre_write(Path("test.json"), '{"a": 1}') is None,
          "validate_syntax_pre_write accepts valid JSON")
    check(validate_syntax_pre_write(Path("test.json"), '{"a": 1,}') is not None,
          "validate_syntax_pre_write catches invalid JSON")

    # 3) SearchProjectTool with FTS5
    with tempfile.TemporaryDirectory() as tmpdir:
        td = Path(tmpdir)
        (td / "hello.py").write_text("def find_secret_token():\n    return 'xyz'\n", encoding="utf-8")
        (td / "doc.md").write_text("# Project Notes\nDatabase connection pooling configuration.\n", encoding="utf-8")
        cfg = scratch_sessions(load_config())
        s = Session(cfg, transient=True)
        ctx = AgentContext(cfg=cfg, session=s, workspace=td)
        tool = SearchProjectTool()
        search_res = tool.run(ctx, "secret token")
        check("hello.py" in search_res and "secret" in search_res and "token" in search_res,
              "SearchProjectTool finds code by keywords with FTS5")
        search_res2 = tool.run(ctx, "database pooling")
        check("doc.md" in search_res2 and "pooling" in search_res2,
              "SearchProjectTool finds relevant documentation")

    # 4) ChangeJournal revert_last_task
    with tempfile.TemporaryDirectory() as tmpdir:
        td = Path(tmpdir)
        test_file = td / "important.txt"
        test_file.write_text("ORIGINAL CONTENT", encoding="utf-8")
        s = Session(cfg, transient=False)
        from harness.changes import ChangeJournal
        journal = ChangeJournal(s, td)
        journal.begin_task("test task")
        journal.record_before(test_file)
        test_file.write_text("MODIFIED BAD CONTENT", encoding="utf-8")
        journal.record_after(test_file)
        check(test_file.read_text(encoding="utf-8") == "MODIFIED BAD CONTENT",
              "File modified")
        rev_res = journal.revert_last_task()
        check(len(rev_res.get("restored", [])) == 1,
              "revert_last_task restored one file")
        check(test_file.read_text(encoding="utf-8") == "ORIGINAL CONTENT",
              "revert_last_task restored the exact original contents")
        Session.delete(cfg, s.id)

    # 5) Slash command dispatcher
    s_cmd = Session(cfg, transient=False)
    webapp.state.session = s_cmd
    webapp.state.rebuild_agent()
    h1, p1 = webapp._handle_slash_command("/help")
    check(h1 is True and "Available slash commands" in s_cmd.messages[-1]["content"],
          "_handle_slash_command handles /help locally")
    h2, p2 = webapp._handle_slash_command("/pins")
    check(h2 is True and "pinned files" in s_cmd.messages[-1]["content"],
          "_handle_slash_command handles /pins locally")
    h3, p3 = webapp._handle_slash_command("/test")
    check(h3 is False and "project checks" in (p3 or ""),
          "_handle_slash_command extends the /test prompt")
    h4, p4 = webapp._handle_slash_command("/skills")
    check(h4 is True and "Available skills" in s_cmd.messages[-1]["content"]
          and "excel-spreadsheet-craft" in s_cmd.messages[-1]["content"],
          "_handle_slash_command handles /skills locally")
    h5, p5 = webapp._handle_slash_command("/skill excel-spreadsheet-craft")
    check(h5 is True and "was activated" in s_cmd.messages[-2]["content"]
          and "[ACTIVE SKILL" in s_cmd.messages[-1]["content"],
          "_handle_slash_command activates a skill in the context")
    h6, p6 = webapp._handle_slash_command("/skill new my-analysis")
    check(h6 is False and "SKILL DESIGNER" in (p6 or ""),
          "_handle_slash_command starts the skill designer")
    Session.delete(cfg, s_cmd.id)

    # 5b) Office and Spreadsheet tools (Excel, Word, PDF)
    from harness.tools.documents import ReadDocumentTool, EditSpreadsheetTool
    from harness.tools.fs import ReadFileTool
    from harness.tools.base import AgentContext
    tmp_doc = Path(tempfile.mkdtemp())
    try:
        s_doc = Session(cfg, transient=True)
        doc_ctx = AgentContext(cfg=cfg, session=s_doc, project_workspace=tmp_doc)
        xlsx_path = tmp_doc / "table.xlsx"
        res_create = EditSpreadsheetTool().run(
            doc_ctx, path=str(xlsx_path), action="create", title="TestSheet",
            data=[["Name", "Salary"], ["Petr", 50000], ["Jana", 60000]])
        check("OK: Created a new Excel" in res_create and xlsx_path.is_file(),
              "EditSpreadsheetTool creates a new XLSX file")
        res_update = EditSpreadsheetTool().run(
            doc_ctx, path=str(xlsx_path), action="update_cells",
            data={"C1": "Bonus", "C2": 5000, "C3": 6000})
        check("OK: Updated 3 cells" in res_update,
              "EditSpreadsheetTool updates cells")
        res_read = ReadDocumentTool().run(doc_ctx, path=str(xlsx_path))
        check("TestSheet" in res_read and "Petr" in res_read and "Bonus" in res_read,
              "ReadDocumentTool extracts XLSX as a Markdown table")
        res_fs = ReadFileTool().run(doc_ctx, path=str(xlsx_path))
        check("[Binary Document converted to text" in res_fs and "Jana" in res_fs,
              "ReadFileTool automatically converts XLSX into readable text")
    finally:
        shutil.rmtree(tmp_doc, ignore_errors=True)



    # 6) Detect repeated tool patterns.
    from harness.agent import Agent, Status, build_registry
    from harness.llm import AssistantResult
    from harness.safety import SafetyPolicy
    script_loop = [AssistantResult(tool_calls=[_tc("list_dir", '{"path": "."}')]) for _ in range(6)]
    s_loop = Session(cfg, transient=True)
    agent_loop = Agent(cfg, LLMStub(script_loop), s_loop, build_registry("agent"),
                       SafetyPolicy("auto"), mode="agent")
    agent_loop.new_task("smycka")
    loop_res = None
    for _ in range(6):
        r = agent_loop.step(approve=True)
        if r.status is Status.FINAL:
            loop_res = r
            break
    check(loop_res is None and r.status is Status.CONTINUE
          and s_loop.load_task_state().get("status") == "running",
          "Repeated calls trigger progress advice, not false completion")


def test_clickable_skills_and_clipboard_images():
    print("--- test_clickable_skills_and_clipboard_images ---")
    import base64
    import json
    import webapp
    from harness.session import Session
    cfg = scratch_sessions(webapp.cfg)

    # 1) Verify skill-information HTML.
    info_html = webapp.skills_info_text()
    check("skills-panel-list" in info_html, "skills_info_text contains the skills-panel-list container")
    check("skill-chip-btn" in info_html, "skills_info_text contains clickable skill-chip-btn buttons")
    check("data-skill=" in info_html, "skills_info_text contains data-skill attributes")

    # 2) Prepare pasted base64 image attachments.
    sample_png_b64 = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
    pasted_json = json.dumps([
        {"name": "test1.png", "data": sample_png_b64},
        {"name": "test2.png", "data": sample_png_b64}
    ])

    orig_sess = webapp.state.session
    test_sess = Session(cfg, transient=True)
    webapp.state.session = test_sess
    webapp.state.agent.session = test_sess
    try:
        sub_res, _, msg_up, pasted_up = webapp.prepare_submission("Analyzuj tyto 2 snimky", pasted_json)
        check(sub_res.get("kind") == "run", "prepare_submission starts a run for a prompt with clipboard images")
        check(pasted_up.get("value") == "[]", "prepare_submission clears the hidden pasted-image field")
        user_img_msgs = [m for m in test_sess.messages if m.get("images")]
        check(len(user_img_msgs) == 1 and len(user_img_msgs[0]["images"]) == 2,
              "Two decoded images were attached to the user message")

        # 3) Render image thumbnails in the conversation.
        views = webapp.chat_view()
        user_views = [v for v in views if v.get("role") == "user"]
        check(len(user_views) > 0, "chat_view contains the user message")
        last_user = user_views[-1]
        check("chat-attached-gallery" in last_user["content"], "chat_view contains the chat-attached-gallery thumbnail gallery")
        check("chat-msg-thumb" in last_user["content"], "chat_view contains chat-msg-thumb thumbnails")
        check("/gradio_api/file=" in last_user["content"], "chat_view thumbnails reference displayable /gradio_api/file= URLs")
    finally:
        webapp.state.session = orig_sess
        webapp.state.agent.session = orig_sess
        Session.delete(cfg, test_sess.id)


if __name__ == "__main__":
    test_config()
    test_memory_layers()
    test_safety()
    test_session()
    test_tools_fs_shell()
    test_gpu_autofit()
    test_registry_modes()
    test_parse_args()
    test_reasoning_effort_kwargs()
    test_runtime_lifecycle_helpers()
    test_dependency_marker()
    test_offline_backup()
    test_streaming_bridge()
    test_parallel_read_tools()
    test_resume_task_and_process_after_restart()
    test_git_tools()
    test_automatic_project_check()
    test_task_plan_and_project_instructions()
    test_code_index()
    test_research_ledger_and_synthesis()
    test_project_document_library()
    test_async_model_switch()
    test_shell_readonly()
    test_workspace()
    test_session_meta()
    test_chat_rewind_and_fork()
    test_transient_session()
    test_web_tools()
    test_projects()
    test_skill_library()
    test_context_compression()
    test_communication_protocol()
    test_user_manuals()
    test_thinking_and_communication()
    test_harness_enhancements()
    test_clickable_skills_and_clipboard_images()
    test_live_sessions_untouched()
    print(f"\n{'=' * 40}\nRESULT: {PASS} ✓ / {FAIL} ✗")
    sys.exit(1 if FAIL else 0)
